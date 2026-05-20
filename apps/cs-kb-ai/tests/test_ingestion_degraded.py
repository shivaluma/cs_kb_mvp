from __future__ import annotations

import io
import unittest

from docx import Document
from openpyxl import Workbook

from app import ingestion, openrouter
from app.schemas import DocumentMetadata, ExtractedUnitsPayload, WorkflowExtractionPayload


class IngestionDegradedDraftTest(unittest.TestCase):
    def setUp(self) -> None:
        self.original_rule_extractor = ingestion.extract_rule_table_units
        self.original_workflow_extractor = ingestion.extract_workflow_units
        self.original_workflow_v3_extractor = ingestion.extract_workflow_units_v3
        self.original_workflow_v2_extractor = ingestion.extract_workflow_units_v2
        self.original_renderer = ingestion.render_pdf_pages_as_data_urls
        self.original_refiner = ingestion.refine_extracted_units
        self.original_source_evidence_formatter = ingestion.format_source_evidence_view
        self.original_metadata_suggester = ingestion.suggest_document_metadata

    def tearDown(self) -> None:
        ingestion.extract_rule_table_units = self.original_rule_extractor
        ingestion.extract_workflow_units = self.original_workflow_extractor
        ingestion.extract_workflow_units_v3 = self.original_workflow_v3_extractor
        ingestion.extract_workflow_units_v2 = self.original_workflow_v2_extractor
        ingestion.render_pdf_pages_as_data_urls = self.original_renderer
        ingestion.refine_extracted_units = self.original_refiner
        ingestion.format_source_evidence_view = self.original_source_evidence_formatter
        ingestion.suggest_document_metadata = self.original_metadata_suggester

    def test_docx_policy_rule_openrouter_disabled_creates_degraded_units(self) -> None:
        ingestion.extract_rule_table_units = lambda _filename, _raw_text: ([], ["openrouter_disabled"])
        data = docx_bytes(
            [
                "Quy định khóa/mở khóa tài khoản",
                "1. Trường hợp khóa tài khoản: Nếu tài khoản đang bị khóa vì lý do khác thì L2 kiểm tra trước khi xử lý.",
                "Lưu ý: không cung cấp thông tin bảo mật cho người không được xác minh.",
            ]
        )

        _raw, _digest, chunks, warnings, enrichment = ingestion.prepare_document_version(
            filename="quy-dinh-khoa-mo-khoa.docx",
            content_type="application/vnd.openxmlformats-officedocument.wordprocessingml.document",
            data=data,
            metadata=DocumentMetadata(owner_team="CS Ops"),
        )

        sections = [chunk["section"] for chunk in chunks]
        self.assertEqual(enrichment["document_type"], "policy_rule")
        self.assertEqual(enrichment["extraction_status"], "degraded")
        self.assertTrue(enrichment["publish_blocked"])
        self.assertIn("full_sop", sections)
        self.assertIn("candidate_rule", sections)
        self.assertIn("candidate_warning", sections)
        self.assertTrue(any("openrouter_disabled" in warning for warning in warnings))
        first_rule = next(chunk for chunk in chunks if chunk["section"] == "candidate_rule")
        self.assertEqual(first_rule["metadata"]["review_status"], "needs_review")
        self.assertEqual(first_rule["metadata"]["source_ref_quality"], "paragraph_only")
        artifacts = enrichment["pipeline_artifacts"]
        self.assertTrue(any(artifact["artifact_type"] == "source_blocks" for artifact in artifacts))
        self.assertTrue(any(artifact["artifact_type"] == "classification_result" for artifact in artifacts))
        self.assertTrue(any(artifact["artifact_type"] == "degraded_draft" for artifact in artifacts))
        self.assertTrue(any(artifact["artifact_type"] == "verification_report" for artifact in artifacts))
        self.assertEqual(enrichment["pipeline_job_status"], "degraded")

    def test_metadata_preview_falls_back_when_openrouter_fails(self) -> None:
        def fail_metadata(*_args, **_kwargs):
            raise RuntimeError("401 Unauthorized")

        ingestion.suggest_document_metadata = fail_metadata

        preview = ingestion.preview_document_metadata(
            filename="refund-policy.txt",
            content_type="text/plain",
            data=(
                "Quy định hoàn tiền cho khách hàng\n"
                "CS kiểm tra thanh toán và tài khoản trước khi xử lý refund."
            ).encode("utf-8"),
        )

        self.assertEqual(preview["title"], "Quy định hoàn tiền cho khách hàng")
        self.assertEqual(preview["suggested_metadata"].source, "metadata_preview_fallback")
        self.assertIn("metadata_preview_fallback_used", preview["suggested_metadata"].extraction_warnings)
        self.assertIn("metadata_preview_fallback_used", preview["signals"]["metadata_warnings"])
        self.assertTrue(any(warning.startswith("openrouter_metadata_suggestion_failed") for warning in preview["signals"]["metadata_warnings"]))
        self.assertIn("refund", preview["suggested_metadata"].tags)

    def test_rule_table_prompt_shape_merges_full_sop_into_units_contract(self) -> None:
        normalized_payload, warnings = openrouter.normalize_rule_table_response_payload(
            {
                "full_sop": {
                    "title": "Quy định làm tròn số tiền",
                    "content": "Quy định cách làm tròn theo từng dòng bảng.",
                    "source_refs": [
                        {"source_type": "docx_table", "source_file": "rounding.docx", "table_index": 0, "row_index": 0}
                    ],
                },
                "units": [
                    {
                        "unit_type": "policy_rule",
                        "title": "beFood bồi hoàn món ăn",
                        "content": "Mốc 500đ: <500 làm tròn xuống, ≥500 làm tròn lên.",
                        "source_refs": [
                            {"source_type": "docx_table", "source_file": "rounding.docx", "table_index": 0, "row_index": 1}
                        ],
                    }
                ],
                "warnings": ["effective_from_missing"],
                "metadata_suggestions": {"sub_type": "financial_threshold_matrix"},
                "coverage_report": {"policy_rule_count": 1},
            }
        )

        payload = ExtractedUnitsPayload.model_validate(normalized_payload)
        self.assertEqual(warnings, ["effective_from_missing"])
        self.assertEqual(payload.units[0].unit_type, "full_sop")
        self.assertEqual(payload.units[0].metadata["retrieval_scope"], "document")
        self.assertEqual(payload.units[0].metadata["metadata_suggestions"]["sub_type"], "financial_threshold_matrix")
        self.assertEqual(payload.units[1].unit_type, "policy_rule")

    def test_excel_text_line_source_refs_are_repaired_to_sheet_context(self) -> None:
        raw_text = "\n".join(
            [
                "# Hotline(chưa áp dụng)",
                "",
                "## Kênh liên hệ: Hotline 1900232345",
                "Ngữ cảnh: Quy định xác minh tài khoản",
                "",
                "# Quy trình xác minh từ 02072025",
                "",
                "## Kênh: Call In App, Non-voice In App, Chat in app, chat Social",
                "Link: https://example.test/cia-chat.pdf",
            ]
        )
        units = [
            {
                "unit_type": "full_sop",
                "title": "Quy định xác minh tài khoản",
                "content": "Workbook xác minh tài khoản.",
                "metadata": {},
                "source_refs": [{"source_type": "text", "source_file": "rules.xlsx", "line_start": 1, "line_end": 9}],
            },
            {
                "unit_type": "related_document",
                "title": "Quy trình xác minh từ 02072025",
                "content": "Link tài liệu liên quan.",
                "metadata": {},
                "source_refs": [{"source_type": "text", "source_file": "rules.xlsx", "line_start": 8, "line_end": 9}],
            },
        ]

        repaired, repair_warnings = openrouter.repair_spreadsheet_source_refs_from_raw_text(
            units,
            "rules.xlsx",
            raw_text,
        )

        self.assertEqual(repair_warnings, ["source_refs_repaired_from_text_lines:2"])
        self.assertEqual(repaired[0]["source_refs"][0]["source_type"], "excel")
        self.assertEqual(repaired[0]["source_refs"][0]["sheet"], "Hotline(chưa áp dụng)")
        self.assertEqual(repaired[1]["source_refs"][0]["sheet"], "Quy trình xác minh từ 02072025")
        self.assertEqual(openrouter.source_ref_validation_errors(repaired, "rules.xlsx"), [])

    def test_ai_breakdown_artifact_is_persisted_for_structuring_attempt(self) -> None:
        def fake_rule_extractor(_filename: str, _raw_text: str):
            openrouter.record_ai_breakdown(
                {
                    "flow": "rule_table",
                    "model": "test/model",
                    "status": "completed",
                    "raw_response": {"text": "{\"units\":[]}", "chars": 12, "truncated": False},
                    "parsed_response": {"json": "{\"units\":[]}", "chars": 12, "truncated": False},
                    "normalized_unit_count": 2,
                    "warnings": [],
                }
            )
            refs = [{"source_type": "docx", "source_file": "email.docx", "paragraph_index": 0}]
            return (
                [
                    {
                        "unit_type": "full_sop",
                        "title": "Quy định xác minh email",
                        "content": "Quy định xử lý email sai định dạng.",
                        "confidence": 0.8,
                        "metadata": {"retrieval_scope": "document"},
                        "source_refs": refs,
                    },
                    {
                        "unit_type": "policy_rule",
                        "title": "Sửa email sai định dạng",
                        "content": "CS sửa các lỗi định dạng email phổ biến trước khi phản hồi.",
                        "confidence": 0.8,
                        "metadata": {"retrieval_scope": "unit"},
                        "source_refs": refs,
                    },
                ],
                ["openrouter_rule_table_extraction_used"],
            )

        ingestion.extract_rule_table_units = fake_rule_extractor

        _raw, _digest, chunks, warnings, enrichment = ingestion.prepare_document_version(
            filename="Quy định xác minh địa chỉ email.docx",
            content_type="application/vnd.openxmlformats-officedocument.wordprocessingml.document",
            data=docx_email_verification_bytes(),
            metadata=DocumentMetadata(owner_team="CS Ops"),
        )

        ai_breakdown = next(artifact for artifact in enrichment["pipeline_artifacts"] if artifact["artifact_type"] == "ai_breakdown")
        self.assertEqual(ai_breakdown["stage"], "ai_structure")
        self.assertEqual(ai_breakdown["payload"]["attempt_count"], 1)
        self.assertEqual(ai_breakdown["payload"]["selected_flow"], "rule_table")
        self.assertEqual(ai_breakdown["payload"]["attempts"][0]["raw_response"]["text"], "{\"units\":[]}")

    def test_source_evidence_view_artifact_is_persisted_for_review_ui(self) -> None:
        def fake_source_formatter(filename: str, raw_text: str, document_type: str, source_type: str):
            openrouter.record_ai_breakdown(
                {
                    "flow": "source_evidence_view",
                    "model": "test/model",
                    "status": "completed",
                    "raw_response": {"text": "{\"markdown\":\"# SOP\"}", "chars": 19, "truncated": False},
                    "parsed_response": {"json": "{\"markdown\":\"# SOP\"}", "chars": 19, "truncated": False},
                    "warnings": [],
                }
            )
            return (
                {
                    "title": filename.rsplit(".", 1)[0],
                    "format": "markdown",
                    "formatter": "ai_source_evidence_view",
                    "model": "test/model",
                    "markdown": "# SOP\n\n- Dòng nguồn đã được format.",
                    "raw_text_chars": len(raw_text),
                    "sections": [{"title": "SOP", "source_hint": "paragraph", "confidence": 0.8}],
                    "warnings": [],
                    "coverage_report": {"raw_text_chars": len(raw_text), "formatted_chars": 32},
                },
                ["openrouter_source_evidence_formatter_used"],
                "",
            )

        ingestion.format_source_evidence_view = fake_source_formatter
        ingestion.extract_rule_table_units = lambda _filename, _raw_text: ([], ["openrouter_invalid_json"])

        _raw, _digest, chunks, warnings, enrichment = ingestion.prepare_document_version(
            filename="Quy định xác minh địa chỉ email.docx",
            content_type="application/vnd.openxmlformats-officedocument.wordprocessingml.document",
            data=docx_email_verification_bytes(),
            metadata=DocumentMetadata(owner_team="CS Ops"),
        )

        source_view = next(
            artifact for artifact in enrichment["pipeline_artifacts"] if artifact["artifact_type"] == "source_evidence_view"
        )
        breakdown = next(
            artifact
            for artifact in enrichment["pipeline_artifacts"]
            if artifact["artifact_type"] == "source_evidence_ai_breakdown"
        )
        self.assertEqual(source_view["stage"], "map")
        self.assertEqual(source_view["status"], "completed")
        self.assertEqual(source_view["payload"]["format"], "markdown")
        self.assertIn("Dòng nguồn đã được format", source_view["payload"]["markdown"])
        self.assertEqual(breakdown["payload"]["selected_flow"], "source_evidence_view")
        unit_types = [chunk["metadata"].get("unit_type") for chunk in chunks]
        self.assertIn("source_evidence_section", unit_types)
        self.assertIn("source_evidence_sections_created", warnings)
        source_section = next(chunk for chunk in chunks if chunk["metadata"].get("unit_type") == "source_evidence_section")
        self.assertEqual(source_section["metadata"]["retrieval_scope"], "source_evidence")
        self.assertTrue(source_section["metadata"]["source_evidence_only"])
        full_sop = next(chunk for chunk in chunks if chunk["metadata"].get("unit_type") == "full_sop")
        self.assertEqual(full_sop["metadata"]["document_layer_role"], "overview")
        self.assertTrue(full_sop["metadata"]["full_source_in_source_evidence_sections"])

    def test_excel_multiple_dated_sheets_creates_candidate_rows_with_scope(self) -> None:
        ingestion.extract_rule_table_units = lambda _filename, _raw_text: ([], ["openrouter_invalid_json"])
        data = workbook_bytes(
            {
                "Từ ngày 01.05.2025": [["Case", "Action"], ["A", "Kiểm tra và xử lý theo quy định"]],
                "Old archive": [["Case", "Action"], ["B", "Không tự động active historical row"]],
            }
        )

        _raw, _digest, chunks, _warnings, enrichment = ingestion.prepare_document_version(
            filename="rules.xlsx",
            content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
            data=data,
            metadata=DocumentMetadata(owner_team="CS Ops"),
        )

        table_rows = [chunk for chunk in chunks if chunk["section"] == "candidate_table_row"]
        self.assertEqual(enrichment["document_type"], "policy_table")
        self.assertEqual(enrichment["extraction_status"], "degraded")
        self.assertEqual(len(table_rows), 2)
        current = next(chunk for chunk in table_rows if chunk["metadata"]["sheet_name"] == "Từ ngày 01.05.2025")
        historical = next(chunk for chunk in table_rows if chunk["metadata"]["sheet_name"] == "Old archive")
        self.assertEqual(current["metadata"]["effective_from"], "2025-05-01")
        self.assertEqual(current["metadata"]["version_scope"], "current_candidate")
        self.assertEqual(historical["metadata"]["version_scope"], "historical_candidate")
        self.assertEqual(current["metadata"]["source_refs"][0]["row_start"], 2)
        self.assertEqual(current["metadata"]["source_refs"][0]["column_names"], ["Case", "Action"])
        self.assertTrue(any(artifact["artifact_type"] == "structuring_plan" for artifact in enrichment["pipeline_artifacts"]))

    def test_excel_hyperlink_rows_create_related_document_candidates(self) -> None:
        ingestion.extract_rule_table_units = lambda _filename, _raw_text: ([], ["openrouter_invalid_json"])
        data = workbook_with_hyperlinks_bytes()

        _raw, _digest, chunks, _warnings, enrichment = ingestion.prepare_document_version(
            filename="Quy định xác minh tài khoản.xlsx",
            content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
            data=data,
            metadata=DocumentMetadata(owner_team="CS Ops"),
        )

        related_rows = [
            chunk for chunk in chunks
            if chunk["metadata"].get("related_documents")
        ]
        self.assertEqual(enrichment["document_type"], "policy_table")
        self.assertGreaterEqual(len(related_rows), 2)
        targets = [
            relation["target_title"]
            for chunk in related_rows
            for relation in chunk["metadata"]["related_documents"]
        ]
        self.assertIn("Call In App, Non-voice In App", targets)
        self.assertIn("Xác minh tài khoản kênh hotline.xlsx", targets)
        self.assertTrue(all(relation.get("source_url", "").startswith("https://") for chunk in related_rows for relation in chunk["metadata"]["related_documents"]))
        self.assertTrue(any(chunk["metadata"].get("hyperlinks") for chunk in related_rows))
        self.assertTrue(all("cell_text" in link for chunk in related_rows for link in chunk["metadata"].get("hyperlinks", [])))

    def test_successful_ai_spreadsheet_structuring_keeps_related_document_units(self) -> None:
        def fake_rule_extractor(_filename: str, _raw_text: str):
            refs = [{"source_type": "excel", "source_file": "rules.xlsx", "sheet": "Rules", "row_start": 1, "row_end": 1}]
            return (
                [
                    {
                        "unit_type": "policy_rule",
                        "title": "Quy định xác minh tài khoản",
                        "content": "CS xác minh theo từng kênh.",
                        "source_refs": refs,
                    }
                ],
                ["openrouter_rule_table_extraction_used"],
            )

        ingestion.extract_rule_table_units = fake_rule_extractor
        data = workbook_with_hyperlinks_bytes()

        _raw, _digest, chunks, _warnings, enrichment = ingestion.prepare_document_version(
            filename="Quy định xác minh tài khoản.xlsx",
            content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
            data=data,
            metadata=DocumentMetadata(owner_team="CS Ops"),
        )

        related_units = [chunk for chunk in chunks if chunk["metadata"].get("unit_type") == "related_document"]
        self.assertEqual(enrichment["extraction_status"], "structured")
        self.assertGreaterEqual(len(related_units), 2)
        self.assertIn("Source URL:", related_units[0]["content"])

    def test_cs_daily_use_workbook_creates_kb_index_plan_and_candidates(self) -> None:
        ingestion.format_source_evidence_view = lambda **_kwargs: ({}, [], "")
        data = cs_daily_use_workbook_bytes()

        _raw, _digest, chunks, warnings, enrichment = ingestion.prepare_document_version(
            filename="CS daily use 2026.xlsx",
            content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
            data=data,
            metadata=DocumentMetadata(owner_team="CS Ops"),
        )

        self.assertEqual(enrichment["document_type"], "kb_index_workbook")
        self.assertEqual(enrichment["source_type"], "excel_workbook")
        self.assertEqual(enrichment["extraction_status"], "structured")
        self.assertIn("kb_index_plan_extraction_used", warnings)

        plan = next(artifact for artifact in enrichment["pipeline_artifacts"] if artifact["artifact_type"] == "kb_index_plan")
        self.assertEqual(plan["stage"], "plan")
        self.assertGreaterEqual(plan["payload"]["summary"]["collection_count"], 5)
        self.assertGreaterEqual(plan["payload"]["summary"]["issue_router_unit_count"], 2)
        self.assertGreaterEqual(plan["payload"]["summary"]["tool_link_count"], 1)
        self.assertGreaterEqual(plan["payload"]["summary"]["unresolved_target_count"], 1)

        unit_types = {chunk["metadata"].get("unit_type") for chunk in chunks}
        self.assertIn("issue_router_unit", unit_types)
        self.assertIn("sop_reference", unit_types)
        self.assertIn("tool_link", unit_types)
        self.assertIn("vip_overlay_rule", unit_types)
        self.assertIn("product_update_note", unit_types)
        self.assertIn("quick_action_rule", unit_types)

        router = next(chunk for chunk in chunks if chunk["metadata"].get("unit_type") == "issue_router_unit")
        self.assertEqual(router["metadata"]["source_refs"][0]["source_type"], "excel")
        self.assertEqual(router["metadata"]["source_refs"][0]["sheet"], "Driver + Rider")
        self.assertEqual(router["metadata"]["source_refs"][0]["row_start"], 2)
        self.assertEqual(router["metadata"]["collection_slug"], "trip-order-issues")
        self.assertIn("requires", {relation["relation_type"] for relation in router["metadata"]["relations"]})

        tool = next(chunk for chunk in chunks if chunk["metadata"].get("unit_type") == "tool_link")
        self.assertEqual(tool["metadata"]["url"], "https://tools.example.com/admin")
        self.assertNotEqual(tool["metadata"].get("unit_type"), "sop_reference")
        self.assertEqual(tool["metadata"]["source_refs"][0]["sheet"], "Link làm việc")

    def test_degraded_policy_text_does_not_duplicate_warning_clone_content(self) -> None:
        ingestion.extract_rule_table_units = lambda _filename, _raw_text: ([], ["openrouter_invalid_json"])

        _raw, _digest, chunks, _warnings, enrichment = ingestion.prepare_document_version(
            filename="Quy định xác minh địa chỉ email.docx",
            content_type="application/vnd.openxmlformats-officedocument.wordprocessingml.document",
            data=docx_email_verification_bytes(),
            metadata=DocumentMetadata(owner_team="CS Ops"),
        )

        self.assertEqual(enrichment["extraction_status"], "degraded")
        atomic_contents = [
            ingestion.content_fingerprint(chunk["content"])
            for chunk in chunks
            if chunk["metadata"].get("unit_type") != "full_sop"
        ]
        self.assertEqual(len(atomic_contents), len(set(atomic_contents)))
        duplicated_title_units = [
            chunk for chunk in chunks
            if chunk["heading"] == "Quy định xác minh địa chỉ email"
            and chunk["metadata"].get("unit_type") in {"candidate_rule", "candidate_warning"}
        ]
        self.assertLessEqual(len(duplicated_title_units), 1)

    def test_unknown_llm_unit_type_is_normalized_before_validation(self) -> None:
        payload = ExtractedUnitsPayload.model_validate(
            {
                "units": [
                    {
                        "unit_type": "procedure_step",
                        "title": "Kiểm tra email",
                        "content": "CS kiểm tra email trên hệ thống.",
                        "source_refs": [{"source_type": "docx", "source_file": "email.docx", "paragraph_index": 1}],
                    }
                ]
            }
        )

        self.assertEqual(payload.units[0].unit_type, "operational_instruction")
        self.assertEqual(payload.units[0].metadata["original_unit_type"], "procedure_step")

    def test_docx_policy_table_creates_table_aware_atomic_units(self) -> None:
        ingestion.extract_rule_table_units = lambda _filename, _raw_text: self.fail("docx policy table should not call OpenRouter")
        data = docx_policy_table_bytes()

        _raw, _digest, chunks, _warnings, enrichment = ingestion.prepare_document_version(
            filename="Quy định làm tròn số tiền.docx",
            content_type="application/vnd.openxmlformats-officedocument.wordprocessingml.document",
            data=data,
            metadata=DocumentMetadata(owner_team="CS Ops"),
        )

        self.assertEqual(enrichment["document_type"], "policy_rule")
        self.assertEqual(enrichment["extraction_status"], "structured")
        unit_types = [chunk["metadata"]["unit_type"] for chunk in chunks]
        self.assertIn("full_sop", unit_types)
        self.assertGreaterEqual(unit_types.count("policy_rule"), 2)
        self.assertGreaterEqual(unit_types.count("exception_rule"), 2)

        befood_rule = next(chunk for chunk in chunks if chunk["heading"] == "beFood - Bồi hoàn liên quan món ăn")
        self.assertEqual(befood_rule["metadata"]["source_ref_quality"], "table_row")
        self.assertEqual(befood_rule["metadata"]["source_refs"][0]["source_type"], "docx_table")
        self.assertIn("cell_text", befood_rule["metadata"]["source_refs"][0])
        self.assertIn("source_cell_text", befood_rule["metadata"])
        self.assertEqual(befood_rule["metadata"]["rounding_threshold"], 500)
        self.assertEqual(len(befood_rule["metadata"]["examples"]), 3)
        self.assertIn("10,450đ -> 10,000đ", befood_rule["content"])

        normal_refund = next(chunk for chunk in chunks if chunk["heading"] == "Dịch vụ khác - Hoàn / rút tiền thông thường")
        self.assertEqual(normal_refund["metadata"]["rounding_threshold"], 300)
        self.assertIn(">300", normal_refund["content"])

        pm04 = next(chunk for chunk in chunks if "PM04" in chunk["heading"])
        self.assertEqual(pm04["metadata"]["unit_type"], "exception_rule")
        self.assertFalse(pm04["metadata"]["rounding_applies"])

        source_blocks = next(artifact for artifact in enrichment["pipeline_artifacts"] if artifact["artifact_type"] == "source_blocks")
        self.assertEqual(source_blocks["payload"]["source_ref_quality"], "table_row")
        self.assertTrue(any(block["type"] == "docx_table_row" for block in source_blocks["payload"]["preview_blocks"]))
        verification = next(artifact for artifact in enrichment["pipeline_artifacts"] if artifact["artifact_type"] == "verification_report")
        self.assertNotIn("missing_atomic_units", verification["payload"]["hard_blockers"])
        self.assertNotIn("degraded_units_require_manual_curation", verification["payload"]["hard_blockers"])

    def test_mixed_docx_communication_guideline_preserves_order_sections_and_risk_units(self) -> None:
        ingestion.extract_rule_table_units = lambda _filename, _raw_text: self.fail("mixed DOCX should not use pure rule_table extraction")
        data = communication_guideline_docx_bytes()

        raw, _digest, chunks, warnings, enrichment = ingestion.prepare_document_version(
            filename="Quy định nội dung phản hồi TX, KH.docx",
            content_type="application/vnd.openxmlformats-officedocument.wordprocessingml.document",
            data=data,
            metadata=DocumentMetadata(owner_team="CS Ops"),
        )

        self.assertEqual(enrichment["document_type"], "policy_rule")
        self.assertEqual(enrichment["sub_type"], "communication_guideline")
        self.assertEqual(enrichment["structure_type"], "mixed_docx")
        self.assertEqual(enrichment["extraction_status"], "structured")
        self.assertIn("mixed_docx_policy_extraction_used", warnings)

        source_blocks = next(artifact for artifact in enrichment["pipeline_artifacts"] if artifact["artifact_type"] == "source_blocks")
        email_open_block = next(
            block for block in source_blocks["payload"]["preview_blocks"]
            if block.get("type") == "docx_table_row" and "Open" in block.get("text", "")
        )
        self.assertEqual(email_open_block["section_path"], ["Mẫu câu chào mở đầu, chào kết và cách xưng hô", "Đối với Email"])
        self.assertLess(raw.index("Open | Xin chào Quý khách hàng"), raw.index("IV. Kiểm soát chất lượng"))

        macro_table = next(chunk for chunk in chunks if chunk["metadata"].get("unit_type") == "macro_table")
        self.assertEqual(macro_table["metadata"]["section_path"], ["Mẫu câu chào mở đầu, chào kết và cách xưng hô", "Đối với Email"])
        self.assertIn("Xin chào Quý khách hàng", "\n".join(macro_table["metadata"]["notes"]))
        self.assertTrue(all(ref["source_type"] == "docx_table" for ref in macro_table["metadata"]["source_refs"]))
        self.assertEqual(
            {row["cells"]["Phần"] for row in macro_table["metadata"]["rows"]},
            {"Open", "Body", "Close"},
        )

        call_chat_units = [
            chunk for chunk in chunks
            if chunk["metadata"].get("block_type") == "list_item"
            and "Đối với Call/Chat" in chunk["metadata"].get("section_path", [])
        ]
        self.assertGreaterEqual(len(call_chat_units), 4)
        self.assertEqual(len({unit["content"] for unit in call_chat_units}), len(call_chat_units))
        self.assertTrue(any("Xin chào anh/chị" in unit["content"] for unit in call_chat_units))
        self.assertTrue(
            all(
                "Đối với Email" not in unit["metadata"].get("section_path", [])[:-1]
                for unit in call_chat_units
            )
        )
        inline_rules = [
            unit for unit in call_chat_units
            if unit["metadata"].get("numbering", {}).get("inline_bullet") is True
        ]
        self.assertEqual(len(inline_rules), 2)
        self.assertTrue(any(unit["metadata"].get("condition") == "Nếu KH/TX chưa đề cập vấn đề" for unit in inline_rules))
        self.assertTrue(any("không hỏi lại" in unit["metadata"].get("action", "") for unit in inline_rules))

        apology_units = [
            chunk for chunk in chunks
            if chunk["metadata"].get("unit_type") == "wording_rule"
        ]
        self.assertTrue(any("xin lỗi" in unit["content"].lower() for unit in apology_units))

        sanction_rule = next(chunk for chunk in chunks if "chế tài" in chunk["content"] and chunk["metadata"].get("unit_type") != "full_sop")
        self.assertEqual(sanction_rule["metadata"]["unit_type"], "compliance_rule")
        self.assertEqual(sanction_rule["metadata"]["risk_level"], "critical")
        self.assertEqual(sanction_rule["metadata"]["actor"], "cs")
        self.assertEqual(sanction_rule["metadata"]["affected_audience"], ["driver", "customer"])
        self.assertNotEqual(sanction_rule["metadata"].get("affected_audience"), ["Customer Service"])

        internal_rule = next(chunk for chunk in chunks if "quy trình xử lý nội bộ" in chunk["content"] and chunk["metadata"].get("unit_type") != "full_sop")
        self.assertEqual(internal_rule["metadata"]["unit_type"], "compliance_rule")
        self.assertEqual(internal_rule["metadata"]["risk_level"], "critical")
        self.assertIn("internal_process_disclosure", internal_rule["metadata"]["risk_category"])

        forbidden_rule = next(chunk for chunk in chunks if "TUYỆT ĐỐI KHÔNG" in chunk["content"])
        self.assertIn(forbidden_rule["metadata"]["risk_level"], {"high", "critical"})

        self.assertTrue(
            all(
                any(
                    ref.get("paragraph_index") is not None
                    or (ref.get("source_type") == "docx_table" and ref.get("row_index") is not None)
                    for ref in chunk["metadata"].get("source_refs", [])
                )
                for chunk in chunks
            )
        )

        verification = next(artifact for artifact in enrichment["pipeline_artifacts"] if artifact["artifact_type"] == "verification_report")
        checks = verification["payload"]["checks"]
        for check in [
            "table_order_preserved",
            "section_path_present",
            "table_rows_have_source_refs",
            "bullet_items_split",
            "email_macro_table_extracted",
            "call_chat_not_nested_under_email",
            "inline_bullets_split",
            "actor_audience_normalized",
            "call_chat_greeting_extracted",
            "apology_wording_rules_extracted",
            "high_risk_forbidden_phrase_detected",
            "critical_internal_disclosure_detected",
            "sanction_disclosure_warning_extracted",
            "internal_process_disclosure_rule_extracted",
        ]:
            self.assertTrue(checks[check], check)

    def test_llm_refine_can_enrich_structured_units_after_extraction(self) -> None:
        ingestion.extract_rule_table_units = lambda _filename, _raw_text: self.fail("docx policy table should not call OpenRouter")

        def fake_refiner(**kwargs):
            refined = []
            for unit in kwargs["units"]:
                metadata = dict(unit.get("metadata") or {})
                metadata["llm_refined"] = True
                metadata["aliases"] = [*metadata.get("aliases", []), "refined lookup alias"]
                refined.append({**unit, "metadata": metadata, "source_refs": unit.get("source_refs") or metadata.get("source_refs")})
            return refined, {"llm_refine_status": "completed", "coverage": {"source": "fake"}}, ["openrouter_refine_used"]

        ingestion.refine_extracted_units = fake_refiner

        _raw, _digest, chunks, warnings, enrichment = ingestion.prepare_document_version(
            filename="Quy định làm tròn số tiền.docx",
            content_type="application/vnd.openxmlformats-officedocument.wordprocessingml.document",
            data=docx_policy_table_bytes(),
            metadata=DocumentMetadata(owner_team="CS Ops"),
        )

        self.assertIn("openrouter_refine_used", warnings)
        delivery_chunks = [chunk for chunk in chunks if chunk["metadata"].get("source_evidence_only") is not True]
        self.assertTrue(all(chunk["metadata"].get("llm_refined") is True for chunk in delivery_chunks))
        self.assertTrue(any("refined lookup alias" in chunk["metadata"].get("aliases", []) for chunk in chunks))
        refinement = next(artifact for artifact in enrichment["pipeline_artifacts"] if artifact["artifact_type"] == "refinement_report")
        self.assertEqual(refinement["payload"]["llm"]["llm_refine_status"], "completed")
        self.assertIn("post_guard", refinement["payload"]["llm"])

    def test_workflow_ai_failure_does_not_create_confirmed_graph(self) -> None:
        classification = type("Classification", (), {"document_type": "workflow_diagram", "source_type": "diagram_pdf", "confidence": 0.78})()
        chunks = ingestion.build_degraded_workflow_draft(
            filename="workflow.pdf",
            raw_text="1. CS tiếp nhận yêu cầu\n2. Nếu đủ thông tin thì xử lý\nLưu ý: SLA 30 phút",
            blocks=[],
            raw_context={},
            classification=classification,
            ai_error="ai_workflow_structuring_failed:invalid_json",
        )

        sections = [chunk.section for chunk in chunks]
        self.assertIn("full_sop", sections)
        self.assertIn("candidate_workflow_text", sections)
        self.assertIn("candidate_step", sections)
        self.assertNotIn("workflow_graph", sections)
        self.assertTrue(all(chunk.metadata["publish_blocked"] for chunk in chunks))
        self.assertEqual(chunks[0].metadata["graph_extraction_status"], "not_reliable_without_layout_review")

    def test_workflow_visual_layout_creates_reviewable_bbox_candidates(self) -> None:
        classification = type("Classification", (), {"document_type": "workflow_diagram", "source_type": "diagram_pdf", "confidence": 0.78})()
        chunks = ingestion.build_degraded_workflow_draft(
            filename="workflow.pdf",
            raw_text="1. CS tiếp nhận yêu cầu",
            blocks=[],
            raw_context={
                "visual_layout": {
                    "summary": {"shape_candidate_count": 1, "edge_candidate_count": 0},
                    "pages": [
                        {
                            "page": 1,
                            "graph_candidate": {
                                "nodes": [
                                    {
                                        "id": "p1_node_1",
                                        "type": "action",
                                        "title": "CS tiếp nhận yêu cầu",
                                        "bbox": [10, 20, 120, 80],
                                    }
                                ],
                                "edge_candidates": [],
                            },
                        }
                    ],
                }
            },
            classification=classification,
            ai_error="ai_workflow_structuring_failed:invalid_json",
        )

        visual_candidate = next(chunk for chunk in chunks if chunk.metadata.get("visual_node_id") == "p1_node_1")
        self.assertEqual(chunks[0].metadata["graph_extraction_status"], "visual_layout_candidates_need_review")
        self.assertEqual(visual_candidate.metadata["source_ref_quality"], "bbox")
        self.assertFalse(visual_candidate.metadata["source_ref_acknowledged"])
        self.assertEqual(visual_candidate.metadata["source_refs"][0]["bbox"], [10, 20, 120, 80])

    def test_workflow_semantic_refine_merges_duplicate_nodes_and_keeps_uncertain_edges(self) -> None:
        classification = type("Classification", (), {"document_type": "workflow_diagram", "source_type": "diagram_pdf", "confidence": 0.78})()
        refinement = ingestion.build_workflow_semantic_refinement(
            filename="workflow.pdf",
            visual_layout=workflow_visual_layout_fixture(),
            source_blocks=[],
            classification=classification,
        )

        self.assertGreaterEqual(refinement["summary"]["deduped_count"], 1)
        graph_nodes = refinement["workflow_graph_candidate"]["nodes"]
        transfer_nodes = [node for node in graph_nodes if "Food Order" in node.get("content", "")]
        self.assertEqual(len(transfer_nodes), 1)
        decision = next(node for node in graph_nodes if node["type"] == "decision")
        self.assertEqual(decision["semantic_node_type"], "decision")
        self.assertTrue(decision["question"].endswith("?"))
        self.assertGreaterEqual(refinement["summary"]["uncertain_edge_count"], 1)
        self.assertEqual(refinement["workflow_graph_candidate"]["topology_source"], "visual_connector_candidates_only")

    def test_workflow_semantic_refine_classifies_annotations_sla_and_audit(self) -> None:
        classification = type("Classification", (), {"document_type": "workflow_diagram", "source_type": "diagram_pdf", "confidence": 0.78})()
        refinement = ingestion.build_workflow_semantic_refinement(
            filename="workflow.pdf",
            visual_layout=workflow_visual_layout_fixture(),
            source_blocks=[],
            classification=classification,
        )

        page = refinement["pages"][0]
        semantic_types = {node["semantic_node_type"] for node in page["semantic_nodes"]}
        annotation_types = {node["semantic_node_type"] for node in page["annotations"]}
        self.assertIn("queue_rule", semantic_types)
        self.assertIn("sla_rule", semantic_types)
        self.assertIn("annotation", annotation_types)
        self.assertIn("audit_rule", annotation_types)
        annotation_ids = {node["id"] for node in page["annotations"]}
        outgoing_from_annotations = [
            edge for edge in [*page["edges"], *page["uncertain_edges"]]
            if edge["from_node"] in annotation_ids
        ]
        self.assertEqual(outgoing_from_annotations, [])

    def test_workflow_semantic_refine_backfills_full_decision_text_from_bboxes(self) -> None:
        classification = type("Classification", (), {"document_type": "workflow_diagram", "source_type": "diagram_pdf", "confidence": 0.78})()
        layout = workflow_visual_layout_fixture()
        layout["pages"][0]["text_blocks"].extend(
            [
                {"id": "p1_text_2a", "page": 1, "text": "2. Thời gian hết hạn", "bbox": [280, 260, 430, 284]},
                {"id": "p1_text_2b", "page": 1, "text": "gửi hình của KH", "bbox": [292, 286, 420, 310]},
                {"id": "p1_text_2c", "page": 1, "text": "trước thời gian hết ca làm", "bbox": [260, 312, 450, 336]},
                {"id": "p1_text_2d", "page": 1, "text": "việc của CS_A ?", "bbox": [292, 338, 420, 362]},
            ]
        )
        refinement = ingestion.build_workflow_semantic_refinement(
            filename="workflow.pdf",
            visual_layout=layout,
            source_blocks=[],
            classification=classification,
        )

        graph_nodes = refinement["workflow_graph_candidate"]["nodes"]
        decision = next(node for node in graph_nodes if "Thời gian hết hạn" in node.get("content", ""))
        self.assertEqual(decision["semantic_node_type"], "decision")
        self.assertIn("gửi hình của KH", decision["content"])
        self.assertEqual(decision["question"], "Thời gian hết hạn gửi hình của KH trước thời gian hết ca làm việc của CS_A?")
        self.assertEqual(refinement["workflow_graph_candidate"]["topology_source"], "visual_connector_candidates_only")

    def test_workflow_semantic_refine_does_not_treat_description_as_script(self) -> None:
        self.assertEqual(
            ingestion.classify_semantic_node_type(
                '3.2. Chuyển case vào\nqueue "Food Order Issue" & note thêm\nthông tin tại mục\ndescription theo quy định',
                "action",
            ),
            "queue_rule",
        )

    def test_workflow_semantic_refine_does_not_infer_edges_from_text_order(self) -> None:
        classification = type("Classification", (), {"document_type": "workflow_diagram", "source_type": "diagram_pdf", "confidence": 0.78})()
        layout = workflow_visual_layout_fixture()
        layout["pages"][0]["graph_candidate"]["edge_candidates"] = []
        refinement = ingestion.build_workflow_semantic_refinement(
            filename="workflow.pdf",
            visual_layout=layout,
            source_blocks=[],
            classification=classification,
        )

        graph = refinement["workflow_graph_candidate"]
        self.assertEqual(graph["edges"], [])
        self.assertEqual(graph["uncertain_edges"], [])
        self.assertEqual(graph["topology_source"], "visual_connector_candidates_only")

    def test_workflow_degraded_fallback_uses_semantic_candidate_types(self) -> None:
        classification = type("Classification", (), {"document_type": "workflow_diagram", "source_type": "diagram_pdf", "confidence": 0.78})()
        visual_layout = workflow_visual_layout_fixture()
        semantic_refinement = ingestion.build_workflow_semantic_refinement(
            filename="workflow.pdf",
            visual_layout=visual_layout,
            source_blocks=[],
            classification=classification,
        )
        chunks = ingestion.build_degraded_workflow_draft(
            filename="workflow.pdf",
            raw_text="1. Hướng dẫn KH cung cấp hình ảnh\n2. KH có gửi hình ảnh?",
            blocks=[],
            raw_context={"visual_layout": visual_layout, "workflow_semantic_refinement": semantic_refinement},
            classification=classification,
            ai_error="ai_workflow_structuring_failed:invalid_json",
        )

        sections = [chunk.section for chunk in chunks]
        self.assertIn("workflow_graph", sections)
        self.assertIn("candidate_action", sections)
        self.assertIn("candidate_decision", sections)
        self.assertIn("candidate_annotation", sections)
        self.assertIn("candidate_sla", sections)
        self.assertIn("candidate_audit_rule", sections)
        self.assertNotIn("candidate_step", sections)
        graph_chunk = next(chunk for chunk in chunks if chunk.section == "workflow_graph")
        self.assertIn("workflow_graph", graph_chunk.metadata)
        self.assertGreater(len(graph_chunk.metadata["workflow_graph"]["nodes"]), 0)
        self.assertGreaterEqual(graph_chunk.metadata["uncertain_edges_count"], 1)
        self.assertTrue(graph_chunk.metadata["topology_review_required"])
        semantic_candidate = next(chunk for chunk in chunks if chunk.section == "candidate_decision")
        self.assertEqual(semantic_candidate.metadata["source_ref_quality"], "bbox")
        self.assertTrue(semantic_candidate.metadata["topology_review_required"])

    def test_workflow_semantic_refine_does_not_treat_numbered_oval_as_start(self) -> None:
        self.assertEqual(
            ingestion.classify_semantic_node_type("10. Tạo case lưu trữ trên hệ thống", "start"),
            "action",
        )
        self.assertEqual(
            ingestion.classify_semantic_node_type("(*) Team Lead sẽ phân quyền tài khoản Pancake", "action"),
            "annotation",
        )
        self.assertEqual(
            ingestion.normalize_decision_question("Yes\nNo\n6. KH/TX cung cấp thông tin\nNo\nYes\n9. KH/TX\nđồng ý với\nkết quả?"),
            "KH/TX đồng ý với kết quả?",
        )

    def test_workflow_ai_graph_is_enriched_with_semantic_node_content(self) -> None:
        semantic_refinement = {
            "workflow_graph_candidate": {
                "nodes": [
                    {
                        "id": "sem_p1_text_2a",
                        "type": "decision",
                        "semantic_node_type": "decision",
                        "title": "2. Thời gian hết hạn",
                        "content": "2. Thời gian hết hạn\ngửi hình của KH\ntrước thời gian hết ca làm\nviệc của CS_A ?",
                        "question": "Thời gian hết hạn gửi hình của KH trước thời gian hết ca làm việc của CS_A?",
                        "actor": "CS_A",
                        "bbox": [280, 260, 450, 362],
                        "source_refs": [{"source_type": "pdf_diagram", "source_file": "workflow.pdf", "page": 1, "bbox": [280, 260, 450, 362]}],
                    }
                ]
            }
        }
        units = [
            {
                "unit_type": "workflow_graph",
                "title": "Graph",
                "content": "Graph summary",
                "metadata": {
                    "workflow_graph": {
                        "workflow_id": "wf",
                        "title": "Graph",
                        "nodes": [
                            {
                                "id": "decision_2",
                                "type": "decision",
                                "title": "2. Thời gian hết hạn",
                                "question": "Thời gian hết hạn?",
                                "content": "",
                            }
                        ],
                        "edges": [],
                    }
                },
            }
        ]

        enriched = ingestion.enrich_workflow_units_with_semantic_refinement(units, semantic_refinement)
        node = enriched[0]["metadata"]["workflow_graph"]["nodes"][0]
        self.assertIn("trước thời gian hết ca làm", node["content"])
        self.assertEqual(node["semantic_node_type"], "decision")
        self.assertEqual(node["bbox"], [280, 260, 450, 362])

    def test_workflow_semantic_validation_flags_orphan_annotations(self) -> None:
        graph = {
            "workflow_id": "test",
            "title": "Test",
            "start_node_id": "start",
            "nodes": [{"id": "start", "type": "start", "title": "Start", "semantic_node_type": "start"}],
            "edges": [],
            "annotations": [{"id": "ann_1", "type": "annotation", "content": "Lưu ý"}],
            "uncertain_edges": [],
            "topology_source": "visual_connector_candidates_only",
        }

        errors = ingestion.validate_semantic_workflow_graph_candidate(graph)
        self.assertIn("orphan_annotation:ann_1", errors)

    def test_workflow_schema_normalizes_null_node_fields(self) -> None:
        payload = WorkflowExtractionPayload.model_validate(
            {
                "document_metadata": {"document_type": "workflow_diagram"},
                "full_sop": {
                    "unit_type": "full_sop",
                    "title": "Workflow",
                    "content": "Full SOP",
                    "source_refs": [{"source_type": "pdf_diagram", "source_file": "workflow.pdf", "page": 1}],
                },
                "workflow_graph": {
                    "workflow_id": "wf",
                    "title": "Workflow",
                    "start_node_id": "node_1",
                    "nodes": [
                        {
                            "id": "node_1",
                            "type": "decision",
                            "title": "KH có gửi hình ảnh?",
                            "content": None,
                            "question": None,
                        }
                    ],
                    "edges": [],
                    "graph_confidence": 0.5,
                },
            }
        )

        self.assertEqual(payload.workflow_graph.nodes[0].content, "")
        self.assertEqual(payload.workflow_graph.nodes[0].question, "KH có gửi hình ảnh?")

    def test_workflow_ai_structuring_receives_visual_context(self) -> None:
        captured: dict[str, object] = {}

        def fake_workflow_extractor(_filename: str, _raw_text: str, page_images=None, visual_context=None):
            captured["page_images"] = page_images
            captured["visual_context"] = visual_context
            return [], ["openrouter_disabled"]

        classification = type("Classification", (), {"document_type": "workflow_diagram", "source_type": "diagram_pdf", "confidence": 0.78})()
        original_workflow_extractor = ingestion.extract_workflow_units
        original_workflow_v3_extractor = ingestion.extract_workflow_units_v3
        original_workflow_v2_extractor = ingestion.extract_workflow_units_v2
        original_render_pdf = ingestion.render_pdf_pages_as_data_urls
        ingestion.extract_workflow_units = fake_workflow_extractor
        ingestion.extract_workflow_units_v3 = lambda _filename, _raw_text, page_images=None, visual_context=None: ([], ["openrouter_disabled"], {})
        ingestion.extract_workflow_units_v2 = lambda _filename, _raw_text, page_images=None, visual_context=None: ([], ["openrouter_disabled"])
        ingestion.render_pdf_pages_as_data_urls = lambda _data: (["data:image/jpeg;base64,abc"], ["pdf_vision_pages_rendered:1"])
        try:
            _chunks, warnings, ai_error = ingestion.try_ai_structuring(
                filename="workflow.pdf",
                content_type="application/pdf",
                data=b"%PDF-1.4",
                raw_text="Quy trình có Yes/No",
                classification=classification,
                visual_layout={
                    "source_type": "pdf_visual_layout",
                    "summary": {"shape_candidate_count": 1, "edge_candidate_count": 1},
                    "pages": [
                        {
                            "page": 1,
                            "image_size": [200, 100],
                            "graph_candidate": {
                                "nodes": [{"id": "p1_node_1", "type": "action", "title": "Start", "bbox": [1, 2, 3, 4]}],
                                "edge_candidates": [],
                            },
                        }
                    ],
                },
            )
        finally:
            ingestion.extract_workflow_units = original_workflow_extractor
            ingestion.extract_workflow_units_v3 = original_workflow_v3_extractor
            ingestion.extract_workflow_units_v2 = original_workflow_v2_extractor
            ingestion.render_pdf_pages_as_data_urls = original_render_pdf

        self.assertIn("visual_graph_context_supplied_to_llm", warnings)
        self.assertIn("openrouter_disabled", ai_error)
        self.assertIsInstance(captured["visual_context"], dict)
        self.assertEqual(captured["visual_context"]["summary"]["shape_candidate_count"], 1)

    def test_workflow_v3_graph_primary_is_used_before_v2_and_legacy_flow(self) -> None:
        captured: dict[str, object] = {}

        def fake_v3_extractor(_filename: str, _raw_text: str, page_images=None, visual_context=None):
            captured["page_images"] = page_images
            captured["visual_context"] = visual_context
            return (
                [
                    {
                        "unit_type": "full_sop",
                        "title": "Quy trình theo dõi case hình ảnh",
                        "content": "CS theo dõi case hình ảnh liên quan món ăn và chuyển MSC theo SLA.",
                        "confidence": 0.82,
                        "metadata": {"retrieval_scope": "document"},
                        "source_refs": [{"source_type": "pdf_diagram", "source_file": "workflow.pdf", "page": 1}],
                    },
                    {
                        "unit_type": "workflow_graph",
                        "title": "Quy trình theo dõi case hình ảnh",
                        "content": "Workflow graph có node và edge.",
                        "confidence": 0.72,
                        "metadata": {
                            "retrieval_scope": "graph",
                            "workflow_graph": {
                                "workflow_id": "wf",
                                "title": "Quy trình theo dõi case hình ảnh",
                                "start_node_id": "start",
                                "nodes": [{"id": "start", "type": "start", "title": "Start"}],
                                "edges": [],
                                "graph_confidence": 0.72,
                            },
                        },
                        "source_refs": [{"source_type": "pdf_diagram", "source_file": "workflow.pdf", "page": 1}],
                    },
                    {
                        "unit_type": "workflow_step",
                        "title": "Hướng dẫn cung cấp hình ảnh",
                        "content": "CS hướng dẫn KH cung cấp hình ảnh theo thời gian quy định.",
                        "confidence": 0.82,
                        "metadata": {"retrieval_scope": "unit"},
                        "source_refs": [{"source_type": "pdf_diagram", "source_file": "workflow.pdf", "page": 1}],
                    },
                ],
                ["openrouter_workflow_v3_extraction_used"],
                {
                    "workflow_fidelity_report": {
                        "strategy": "workflow_v3_graph_primary",
                        "fidelity_score": 0.82,
                        "blockers": [],
                    }
                },
            )

        classification = type("Classification", (), {"document_type": "workflow_diagram", "source_type": "diagram_pdf", "confidence": 0.78})()
        ingestion.extract_workflow_units_v3 = fake_v3_extractor
        ingestion.extract_workflow_units_v2 = lambda *_args, **_kwargs: self.fail("v2 workflow extractor should not run after v3 success")
        ingestion.extract_workflow_units = lambda *_args, **_kwargs: self.fail("legacy workflow extractor should not run after v3 success")
        ingestion.render_pdf_pages_as_data_urls = lambda _data: (["data:image/jpeg;base64,abc"], ["pdf_vision_pages_rendered:1"])

        chunks, warnings, ai_error = ingestion.try_ai_structuring(
            filename="workflow.pdf",
            content_type="application/pdf",
            data=b"%PDF-1.4",
            raw_text="Quy trình có diagram",
            classification=classification,
            visual_layout={"summary": {"shape_candidate_count": 99}},
        )

        self.assertEqual(ai_error, "")
        self.assertIn("workflow_extraction_flow:v3_graph_primary", warnings)
        self.assertIsInstance(captured["visual_context"], dict)
        self.assertGreaterEqual(len(chunks), 3)
        self.assertFalse(any(chunk.metadata.get("extraction_status") == "degraded" for chunk in chunks))

    def test_workflow_v2_vision_primary_is_used_after_v3_failure_before_legacy_flow(self) -> None:
        captured: dict[str, object] = {}

        def fake_v2_extractor(_filename: str, _raw_text: str, page_images=None, visual_context=None):
            captured["page_images"] = page_images
            captured["visual_context"] = visual_context
            return (
                [
                    {
                        "unit_type": "full_sop",
                        "title": "Quy trình theo dõi case hình ảnh",
                        "content": "CS theo dõi case hình ảnh liên quan món ăn và chuyển MSC theo SLA.",
                        "confidence": 0.82,
                        "metadata": {"retrieval_scope": "document"},
                        "source_refs": [{"source_type": "pdf_diagram", "source_file": "workflow.pdf", "page": 1}],
                    },
                    {
                        "unit_type": "workflow_graph",
                        "title": "Quy trình theo dõi case hình ảnh",
                        "content": "Workflow graph có node và edge.",
                        "confidence": 0.72,
                        "metadata": {
                            "retrieval_scope": "graph",
                            "workflow_graph": {
                                "workflow_id": "wf",
                                "title": "Quy trình theo dõi case hình ảnh",
                                "start_node_id": "start",
                                "nodes": [{"id": "start", "type": "start", "title": "Start"}],
                                "edges": [],
                                "graph_confidence": 0.72,
                            },
                        },
                        "source_refs": [{"source_type": "pdf_diagram", "source_file": "workflow.pdf", "page": 1}],
                    },
                    {
                        "unit_type": "workflow_step",
                        "title": "Hướng dẫn cung cấp hình ảnh",
                        "content": "CS hướng dẫn KH cung cấp hình ảnh theo thời gian quy định.",
                        "confidence": 0.82,
                        "metadata": {"retrieval_scope": "unit"},
                        "source_refs": [{"source_type": "pdf_diagram", "source_file": "workflow.pdf", "page": 1}],
                    },
                ],
                ["openrouter_workflow_v2_extraction_used"],
            )

        classification = type("Classification", (), {"document_type": "workflow_diagram", "source_type": "diagram_pdf", "confidence": 0.78})()
        ingestion.extract_workflow_units_v3 = lambda _filename, _raw_text, page_images=None, visual_context=None: (
            [],
            ["workflow_v3_fidelity_failed:workflow_v3_missing_visible_steps:1"],
            {"workflow_fidelity_report": {"blockers": ["workflow_v3_missing_visible_steps:1"]}},
        )
        ingestion.extract_workflow_units_v2 = fake_v2_extractor
        ingestion.extract_workflow_units = lambda *_args, **_kwargs: self.fail("legacy workflow extractor should not run after v2 success")
        ingestion.render_pdf_pages_as_data_urls = lambda _data: (["data:image/jpeg;base64,abc"], ["pdf_vision_pages_rendered:1"])

        chunks, warnings, ai_error = ingestion.try_ai_structuring(
            filename="workflow.pdf",
            content_type="application/pdf",
            data=b"%PDF-1.4",
            raw_text="Quy trình có diagram",
            classification=classification,
            visual_layout={"summary": {"shape_candidate_count": 99}},
        )

        self.assertEqual(ai_error, "")
        self.assertIn("workflow_extraction_flow:v2_vision_primary", warnings)
        self.assertIsInstance(captured["visual_context"], dict)
        self.assertGreaterEqual(len(chunks), 3)
        self.assertFalse(any(chunk.metadata.get("extraction_status") == "degraded" for chunk in chunks))

    def test_workflow_ai_missing_required_layers_is_not_structured_success(self) -> None:
        def fake_workflow_extractor(_filename: str, _raw_text: str, page_images=None, visual_context=None):
            return (
                [
                    {
                        "unit_type": "full_sop",
                        "title": "Quy trình cần review",
                        "content": "Model không trả full_sop. Backend giữ bản nháp này để CS Ops review lại từ source.",
                        "confidence": 0.45,
                        "metadata": {"retrieval_scope": "document"},
                    },
                    {
                        "unit_type": "workflow_graph",
                        "title": "Quy trình cần review",
                        "content": "Workflow graph tối thiểu.",
                        "confidence": 0.25,
                        "metadata": {"retrieval_scope": "graph", "workflow_graph": {"nodes": [], "edges": []}},
                    },
                ],
                [
                    "full_sop_missing_from_model_synthesized_for_review",
                    "workflow_graph_missing_from_model_synthesized_for_review",
                ],
            )

        classification = type("Classification", (), {"document_type": "workflow_diagram", "source_type": "diagram_pdf", "confidence": 0.78})()
        ingestion.extract_workflow_units_v2 = lambda _filename, _raw_text, page_images=None, visual_context=None: ([], ["openrouter_disabled"])
        ingestion.extract_workflow_units = fake_workflow_extractor
        ingestion.render_pdf_pages_as_data_urls = lambda _data: (["data:image/jpeg;base64,abc"], [])

        chunks, warnings, ai_error = ingestion.try_ai_structuring(
            filename="workflow.pdf",
            content_type="application/pdf",
            data=b"%PDF-1.4",
            raw_text="1. CS tiếp nhận yêu cầu\n2. Nếu đủ thông tin thì xử lý",
            classification=classification,
            visual_layout={},
        )

        self.assertEqual(chunks, [])
        self.assertIn("full_sop_missing_from_model_synthesized_for_review", warnings)
        self.assertIn("workflow_graph_missing_from_model_synthesized_for_review", ai_error)

    def test_workflow_ai_without_atomic_units_is_not_structured_success(self) -> None:
        classification = type("Classification", (), {"document_type": "workflow_diagram", "source_type": "diagram_pdf", "confidence": 0.78})()

        def fake_workflow_extractor(_filename: str, _raw_text: str, page_images=None, visual_context=None):
            return (
                [
                    {
                        "unit_type": "full_sop",
                        "title": "Quy trình",
                        "content": "Nội dung tổng quan.",
                        "confidence": 0.8,
                        "metadata": {"retrieval_scope": "document"},
                    },
                    {
                        "unit_type": "workflow_graph",
                        "title": "Quy trình",
                        "content": "Graph có node.",
                        "confidence": 0.75,
                        "metadata": {"retrieval_scope": "graph", "workflow_graph": {"nodes": [{"id": "start"}], "edges": []}},
                    },
                ],
                [],
            )

        ingestion.extract_workflow_units = fake_workflow_extractor
        ingestion.extract_workflow_units_v2 = lambda _filename, _raw_text, page_images=None, visual_context=None: ([], ["openrouter_disabled"])
        ingestion.render_pdf_pages_as_data_urls = lambda _data: (["data:image/jpeg;base64,abc"], [])
        chunks, warnings, ai_error = ingestion.try_ai_structuring(
            filename="workflow.pdf",
            content_type="application/pdf",
            data=b"%PDF-1.4",
            raw_text="1. CS tiếp nhận yêu cầu",
            classification=classification,
            visual_layout={},
        )

        self.assertEqual(chunks, [])
        self.assertIn("openrouter_disabled", warnings)
        self.assertIn("missing_atomic_workflow_units", ai_error)

    def test_successful_ai_extraction_creates_structured_draft(self) -> None:
        ingestion.extract_rule_table_units = lambda _filename, _raw_text: (
            [
                {
                    "unit_type": "full_sop",
                    "title": "Quy định xử lý",
                    "content": "Full SOP content",
                    "confidence": 0.92,
                    "metadata": {"retrieval_scope": "document"},
                    "source_refs": [{"source_type": "docx", "source_file": "policy.docx", "paragraph_index": 0}],
                },
                {
                    "unit_type": "policy_rule",
                    "title": "Kiểm tra điều kiện",
                    "content": "Nếu có điều kiện A thì xử lý theo nguồn.",
                    "confidence": 0.91,
                    "metadata": {},
                    "source_refs": [{"source_type": "docx", "source_file": "policy.docx", "paragraph_index": 1}],
                },
            ],
            [],
        )
        data = docx_bytes(["Quy định xử lý", "Nếu có điều kiện A thì xử lý theo nguồn."])

        _raw, _digest, chunks, _warnings, enrichment = ingestion.prepare_document_version(
            filename="policy.docx",
            content_type="application/vnd.openxmlformats-officedocument.wordprocessingml.document",
            data=data,
            metadata=DocumentMetadata(owner_team="CS Ops"),
        )

        self.assertEqual(enrichment["extraction_status"], "structured")
        self.assertFalse(enrichment["publish_blocked"])
        self.assertIn("policy_rule", [chunk["section"] for chunk in chunks])
        self.assertTrue(all(chunk["metadata"]["extraction_status"] == "structured" for chunk in chunks))
        self.assertTrue(all(chunk["metadata"]["index_eligible"] is False for chunk in chunks))
        self.assertTrue(any(artifact["artifact_type"] == "ai_structured_payload" for artifact in enrichment["pipeline_artifacts"]))


def docx_bytes(paragraphs: list[str]) -> bytes:
    document = Document()
    for paragraph in paragraphs:
        document.add_paragraph(paragraph)
    buffer = io.BytesIO()
    document.save(buffer)
    return buffer.getvalue()


def docx_policy_table_bytes() -> bytes:
    document = Document()
    document.add_paragraph("QUY ĐỊNH LÀM TRÒN SỐ TIỀN")
    headers = ["Dịch vụ", "Trường hợp", "Quy tắc làm tròn", "Lưu ý"]
    rows = [
        [
            "beFood",
            "Bồi hoàn liên quan món ăn",
            "Mốc 500đ: <500↓, ≥500↑",
            "Làm tròn giá trị bồi hoàn cuối cùng\nví dụ:\n10,450đ -> 10,000 đ\n10,500đ (hoặc 10,560đ) -> 11,000 đ",
        ],
        ["beFood", "Các trường hợp khác (begin-end...)", "Không áp dụng", "Làm tròn giá trị bồi hoàn cuối cùng"],
        [
            "Dịch vụ khác",
            "Hoàn / rút tiền thông thường",
            "Mốc: 300đ: >300↑, ≤300↓",
            "Tại mỗi bước tính tiền, sau khi tính ra số tiền CS cần phải làm tròn theo quy tắc trước khi thực hiện bước tiếp theo.\nVD:\n10,400 -> 11000\n10,200 -> 10,000",
        ],
        ["Dịch vụ khác", "Hoàn / rút chiết khấu ĐT", "Không áp dụng", "Tại mỗi bước tính tiền, sau khi tính ra số tiền CS cần phải làm tròn theo quy tắc trước khi thực hiện bước tiếp theo."],
        ["Dịch vụ khác", "Hoàn KH về PTTT đã dùng (PM04)", "Không áp dụng", "Tại mỗi bước tính tiền, sau khi tính ra số tiền CS cần phải làm tròn theo quy tắc trước khi thực hiện bước tiếp theo."],
    ]
    table = document.add_table(rows=1, cols=len(headers))
    for index, header in enumerate(headers):
        table.rows[0].cells[index].text = header
    for row in rows:
        cells = table.add_row().cells
        for index, value in enumerate(row):
            cells[index].text = value
    buffer = io.BytesIO()
    document.save(buffer)
    return buffer.getvalue()


def communication_guideline_docx_bytes() -> bytes:
    document = Document()
    document.add_paragraph("Quy định nội dung phản hồi TX, KH")
    document.add_paragraph("Mẫu câu chào mở đầu, chào kết và cách xưng hô")
    document.add_paragraph("Đối với Email")
    table = document.add_table(rows=1, cols=2)
    table.rows[0].cells[0].text = "Phần"
    table.rows[0].cells[1].text = "Nội dung"
    for label, text in [
        ("Open", "Xin chào Quý khách hàng,"),
        ("Body", "CS phản hồi theo đúng nội dung xử lý trong case và không tự thêm thông tin ngoài SOP."),
        ("Close", "Trân trọng cảm ơn Quý khách hàng đã liên hệ be."),
    ]:
        cells = table.add_row().cells
        cells[0].text = label
        cells[1].text = text
    document.add_paragraph('Lưu ý: Câu "Xin chào Quý khách hàng" chỉ áp dụng cho Email gửi KH.')
    document.add_paragraph("Đối với Call/Chat")
    document.add_paragraph("CS chào KH/TX: Xin chào anh/chị, em là CS be đang hỗ trợ mình.", style="List Bullet")
    document.add_paragraph("CS xác nhận đúng vấn đề của KH/TX trước khi phản hồi hướng xử lý.", style="List Bullet")
    document.add_paragraph(
        "Câu dẫn vào nội dung hỗ trợ sau khi chào: "
        "- Nếu KH/TX chưa đề cập vấn đề: CS khai thác vấn đề. "
        "- Nếu KH/TX đã đề cập vấn đề: CS không hỏi lại, chỉ xác nhận hoặc hỏi thêm nếu chưa rõ."
    )
    document.add_paragraph("II. Quy định wording xin lỗi")
    document.add_paragraph("Chỉ dùng từ xin lỗi khi lỗi phát sinh từ phía be hoặc trải nghiệm dịch vụ bị ảnh hưởng.", style="List Bullet")
    document.add_paragraph("Không lạm dụng xin lỗi trong mọi câu phản hồi; ưu tiên ghi nhận thông tin và hướng xử lý.", style="List Bullet")
    document.add_paragraph("III. Nội dung không được cung cấp")
    document.add_paragraph("Không cung cấp ngưỡng/lý do chế tài cho TX/KH trong phản hồi.", style="List Bullet")
    document.add_paragraph("Không chủ động cung cấp quy trình xử lý nội bộ cho TX/KH.", style="List Bullet")
    document.add_paragraph("IV. Kiểm soát chất lượng")
    document.add_paragraph('TUYỆT ĐỐI KHÔNG NÓI "Đóng hỗ trợ tại đây" vì sẽ bị chấm lỗi nội dung phản hồi.')
    document.add_paragraph("CS có thể bị chấm lỗi nếu tự ý cung cấp nội dung nội bộ hoặc sai mẫu phản hồi.")
    buffer = io.BytesIO()
    document.save(buffer)
    return buffer.getvalue()


def docx_email_verification_bytes() -> bytes:
    return docx_bytes(
        [
            "Quy định xác minh địa chỉ email",
            "1.Đối với các email có định dạng: gmai.com, gmal.com, gmil.com.com, gmail.con, gamil.com, gmall.com, gmaol.com, outlook.con, @hotmail.con, @yahoo.con, @gmaik.",
            "CS chủ động sửa thành định dạng mail đúng và phản hồi cho khách hàng/tài xế, không cần liên hệ xác minh địa chỉ mail. Trường hợp không sửa mail nhưng vẫn gửi mail phản hồi => lỗi ZT",
            "2. Đối với các email có định dạng sai khác",
            "Lấy SĐT/user ID ở mục contact information kiểm tra trên hệ thống Bizops",
            "Có mail đúng định dạng: soạn nội dung và gửi đến địa chỉ mail trên hệ thống Admin",
            "Không có mail đúng định dạng: gửi phản hồi đến email đang có => Resolve case.",
            "Trường hợp KH/TX liên hệ lại khiếu nại không nhận được email, CS check case liên quan nếu thấy đã gửi mail cho KH/TX, báo KH/TX be đã gửi mail vào [địa chỉ mail KH/TX đã đăng kí trên app], nếu KH/TX báo mail đó lỗi, CS cung cấp kết quả theo mail đã gửi và hướng dẫn KH tự thay đổi email trên app",
            "Nếu là TX thì CS thông tin cho TX việc mail trên hệ thống đang không chính xác => thực hiện quy trình đổi mail cho TX",
            "Lưu ý:",
            "Không gửi mail theo quy trình: lỗi ZT",
        ]
    )


def workflow_visual_layout_fixture() -> dict:
    return {
        "filename": "workflow.pdf",
        "source_type": "pdf_visual_layout",
        "summary": {
            "page_count": 1,
            "shape_candidate_count": 7,
            "connector_candidate_count": 2,
            "edge_candidate_count": 2,
            "confidence": 0.62,
        },
        "pages": [
            {
                "page": 1,
                "image_size": [1200, 900],
                "text_blocks": [
                    {
                        "id": "p1_text_audit",
                        "page": 1,
                        "text": "Quy định audit:\n- CS_A chuyển trễ: ZT",
                        "bbox": [700, 500, 980, 580],
                    }
                ],
                "graph_candidate": {
                    "nodes": [
                        {
                            "id": "p1_node_1",
                            "page": 1,
                            "type": "start",
                            "title": "KH khiếu nại các vấn đề liên quan đến đơn hàng beFood",
                            "bbox": [40, 80, 220, 170],
                            "confidence": 0.68,
                        },
                        {
                            "id": "p1_node_2",
                            "page": 1,
                            "type": "decision",
                            "title": "3. KH có gửi\nhình ảnh?",
                            "bbox": [280, 80, 420, 190],
                            "confidence": 0.76,
                        },
                        {
                            "id": "p1_node_3",
                            "page": 1,
                            "type": "action",
                            "title": "3.2. Chuyển case vào\nqueue \"Food Order",
                            "bbox": [480, 95, 610, 230],
                            "confidence": 0.68,
                        },
                        {
                            "id": "p1_node_4",
                            "page": 1,
                            "type": "action",
                            "title": "3.2. Chuyển case vào\nqueue \"Food Order Issue\" & note thêm thông tin",
                            "bbox": [560, 95, 740, 230],
                            "confidence": 0.68,
                        },
                        {
                            "id": "p1_node_5",
                            "page": 1,
                            "type": "action",
                            "title": "(a) Quy định note description: tóm tắt vấn đề_thời hạn hết hạn gửi hình của KH",
                            "bbox": [760, 100, 980, 220],
                            "confidence": 0.68,
                        },
                        {
                            "id": "p1_node_6",
                            "page": 1,
                            "type": "action",
                            "title": "2.1. Theo dõi case và xử lý bước tiếp theo",
                            "bbox": [260, 320, 450, 430],
                            "confidence": 0.68,
                        },
                        {
                            "id": "p1_node_7",
                            "page": 1,
                            "type": "action",
                            "title": "9.2 CS_B thực hiện bước tiếp theo và đảm bảo case chuyển MSC trễ nhất là 30 phút",
                            "bbox": [480, 320, 780, 430],
                            "confidence": 0.68,
                        },
                        {
                            "id": "p1_node_8",
                            "page": 1,
                            "type": "end",
                            "title": "End",
                            "bbox": [1000, 320, 1120, 390],
                            "confidence": 0.68,
                        },
                    ],
                    "edge_candidates": [
                        {
                            "id": "p1_connector_1",
                            "page": 1,
                            "from_node": "p1_node_2",
                            "to_node": "p1_node_4",
                            "condition": "yes",
                            "bbox": [420, 140, 560, 140],
                            "confidence": 0.54,
                            "direction_reason": "left_to_right_geometric_guess",
                            "review_status": "needs_review",
                        },
                        {
                            "id": "p1_connector_2",
                            "page": 1,
                            "from_node": "p1_node_4",
                            "to_node": "p1_node_6",
                            "condition": "next",
                            "bbox": [650, 230, 650, 320],
                            "confidence": 0.5,
                            "direction_reason": "top_to_bottom_geometric_guess",
                            "review_status": "needs_review",
                        },
                    ],
                    "graph_confidence": 0.55,
                    "requires_human_review": True,
                    "review_reason": "Needs review",
                },
            }
        ],
    }


def workbook_bytes(sheets: dict[str, list[list[str]]]) -> bytes:
    workbook = Workbook()
    first = True
    for sheet_name, rows in sheets.items():
        worksheet = workbook.active if first else workbook.create_sheet()
        first = False
        worksheet.title = sheet_name
        for row in rows:
            worksheet.append(row)
    buffer = io.BytesIO()
    workbook.save(buffer)
    return buffer.getvalue()


def workbook_with_hyperlinks_bytes() -> bytes:
    workbook = Workbook()
    worksheet = workbook.active
    worksheet.title = "Quy trình xác minh từ 02072025"
    worksheet.append(["Kênh", "Link"])
    worksheet.append(["Call In App, Non-voice In App", "Link"])
    worksheet["B2"].hyperlink = "https://example.com/call-in-app.pdf"
    worksheet.append(["Hotline 1900232345", "Link"])
    worksheet["B3"].hyperlink = "https://example.com/hotline.xlsx"
    worksheet["B3"].hyperlink.display = "Xác minh tài khoản kênh hotline.xlsx"
    buffer = io.BytesIO()
    workbook.save(buffer)
    return buffer.getvalue()


def cs_daily_use_workbook_bytes() -> bytes:
    workbook = Workbook()
    workbook.remove(workbook.active)

    sheets = {
        "Overal": [
            ["Quy định/quy trình xử lý", "Số lượng"],
            ["Quy định làm việc CCU,PCU", "2"],
            ["Link làm việc", "1"],
        ],
        "Quy định làm việc CCU,PCU": [
            ["STT", "Tên quy trình/quy định", "Nội dung", "Note"],
            ["1", "Quy định sử dụng tasklist", "CS tạo tasklist theo đúng quy định", ""],
        ],
        "Quy định chung": [
            ["STT", "Vấn đề", "SOP"],
            ["1", "Xác minh tài khoản", "Quy định xác minh tài khoản"],
        ],
        "Driver + Rider": [
            ["STT", "Vấn đề", "Service", "SOP", "File làm việc liên quan", "Note"],
            ["1", "Tạo tasklist Tech khi app lỗi", "beFood", "Quy định sử dụng tasklist", "Link", "CS tạo tasklist và ping Tech"],
        ],
        "Driver+Cleaner": [
            ["STT", "Vấn đề", "Service", "SOP", "File làm việc liên quan", "Note"],
            ["1", "Chuyển queue khi cần cleaner xử lý", "Delivery", "Quy trình chuyển queue Cleaner", "Link", ""],
        ],
        "Rider": [
            ["Vấn đề", "Service", "SOP", "File làm việc liên quan", "Note"],
            ["Khách hỏi hoàn tiền", "Payment", "Quy định hoàn tiền", "Link", ""],
        ],
        "Cleaner": [
            ["Vấn đề", "Service", "SOP", "File làm việc liên quan", "Note"],
            ["Cleaner báo không nhận job", "Delivery", "Quy trình Cleaner", "Link", ""],
        ],
        "MCU": [
            ["Vấn đề", "SOP"],
            ["Merchant cần cập nhật menu", "Quy trình MCU cập nhật menu"],
        ],
        "Link làm việc": [
            ["Nhóm", "Tên file/ hệ thống", "Link", "Note"],
            ["Admin", "CS Admin Tool", "https://tools.example.com/admin", "Mở admin để kiểm tra tài khoản"],
        ],
        "VIP": [
            ["Chủ đề", "Chi tiết vấn đề", "Điểm khác với KH normal", "Link SOP"],
            ["VIP refund", "KH VIP yêu cầu hoàn tiền", "Cần lead review", "Quy định VIP refund"],
        ],
        "Tính năng sản phẩm mới": [
            ["Đối tượng", "Tính năng sản phẩm mới", "Nội dung chi tiết"],
            ["Rider", "Tính năng đặt món mới", "CS tham khảo SOP launch feature"],
        ],
    }

    for sheet_name, rows in sheets.items():
        worksheet = workbook.create_sheet(sheet_name)
        for row in rows:
            worksheet.append(row)
    worksheet = workbook["Driver + Rider"]
    worksheet["E2"].hyperlink = "https://docs.example.com/tasklist"
    worksheet["E2"].hyperlink.display = "Quy định sử dụng tasklist"
    worksheet = workbook["Rider"]
    worksheet["D2"].hyperlink = "https://docs.example.com/refund"
    worksheet["D2"].hyperlink.display = "Quy định hoàn tiền"

    buffer = io.BytesIO()
    workbook.save(buffer)
    return buffer.getvalue()


if __name__ == "__main__":
    unittest.main()
