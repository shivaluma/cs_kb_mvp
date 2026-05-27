from __future__ import annotations

import unittest

from app.chunking.evidence_bound_chunker import build_evidence_bound_chunks, legacy_chunks_from_evidence
from app.ingestion import extract_raw_evidence, parse_document_blocks, prepare_document_version, verification_report_payload
from app.ir.document_evidence_graph import build_document_evidence_graph
from app.ir.schema import SemanticUnit
from app.schemas import DocumentMetadata
from app.validators.chunk_support import validate_chunks_supported
from app.validators.evidence import validate_semantic_units_supported
from app.validators.workflow_graph import validate_workflow_graph


class EvidenceGraphIngestionTest(unittest.TestCase):
    def test_text_evidence_graph_preserves_line_source_elements(self) -> None:
        raw_text = "Greeting SOP\nCS greets the customer.\nCS confirms the issue."
        blocks = parse_document_blocks("greeting.txt", "text/plain", raw_text, {})

        graph = build_document_evidence_graph(
            filename="greeting.txt",
            content_type="text/plain",
            data=raw_text.encode("utf-8"),
            raw_text=raw_text,
            raw_context={},
            blocks=blocks,
            classification=None,
        )

        text_elements = [element for element in graph.source_elements if element.kind == "text"]
        self.assertEqual(len(text_elements), 3)
        self.assertEqual(text_elements[0].provenance.source_locator["line_start"], 1)
        self.assertEqual(text_elements[2].provenance.source_locator["line_end"], 3)
        self.assertTrue(text_elements[0].provenance.raw_excerpt_hash)

    def test_non_ai_fallback_chunks_are_evidence_bound(self) -> None:
        raw_text = "Greeting SOP\nCS greets the customer.\nCS confirms the issue."
        _raw, _digest, chunks, _warnings, enrichment = prepare_document_version(
            filename="greeting.txt",
            content_type="text/plain",
            data=raw_text.encode("utf-8"),
            metadata=DocumentMetadata(owner_team="CS Ops"),
        )

        self.assertEqual(enrichment["document_type"], "text_sop")
        self.assertIn("document_evidence_graph", {artifact["artifact_type"] for artifact in enrichment["pipeline_artifacts"]})
        draft_chunks = [chunk for chunk in chunks if chunk["metadata"].get("unit_type") != "source_evidence_section"]
        self.assertTrue(draft_chunks)
        for chunk in draft_chunks:
            metadata = chunk["metadata"]
            self.assertTrue(metadata.get("source_element_ids"), chunk)
            self.assertTrue(metadata.get("evidence_hash"), chunk)
            self.assertTrue(metadata.get("source_refs"), chunk)

    def test_prepare_document_version_records_durable_map_unit_extracts(self) -> None:
        raw_text = "Greeting SOP\nCS greets the customer.\nCS confirms the issue."
        _raw, _digest, _chunks, _warnings, enrichment = prepare_document_version(
            filename="greeting.txt",
            content_type="text/plain",
            data=raw_text.encode("utf-8"),
            metadata=DocumentMetadata(owner_team="CS Ops"),
        )

        artifact = next(
            artifact
            for artifact in enrichment["pipeline_artifacts"]
            if artifact["artifact_type"] == "map_unit_extracts"
        )

        self.assertEqual(artifact["stage"], "map")
        self.assertEqual(artifact["payload"]["source_element_count"], 3)
        self.assertEqual(artifact["payload"]["unit_count"], 3)
        self.assertEqual(artifact["payload"]["mapped_source_element_count"], 3)
        self.assertEqual(artifact["payload"]["unmapped_source_element_count"], 0)
        self.assertTrue(artifact["payload"]["units"])
        first_unit = artifact["payload"]["units"][0]
        self.assertEqual(first_unit["status"], "completed")
        self.assertEqual(first_unit["source_element_ids"], ["element_1"])
        self.assertTrue(first_unit["evidence_hash"])
        self.assertEqual(first_unit["source_refs"][0]["line_start"], 1)

    def test_prepare_document_version_records_reviewable_compilation_plan(self) -> None:
        raw_text = "Greeting SOP\nCS greets the customer.\nCS confirms the issue."
        _raw, _digest, _chunks, _warnings, enrichment = prepare_document_version(
            filename="greeting.txt",
            content_type="text/plain",
            data=raw_text.encode("utf-8"),
            metadata=DocumentMetadata(owner_team="CS Ops"),
        )

        artifact = next(
            artifact
            for artifact in enrichment["pipeline_artifacts"]
            if artifact["artifact_type"] == "document_compilation_plan"
        )

        payload = artifact["payload"]
        self.assertEqual(artifact["stage"], "plan")
        self.assertEqual(payload["plan_version"], "document_compilation_plan_v1")
        self.assertEqual(payload["document_type"], "text_sop")
        self.assertTrue(payload["human_approval_required"])
        self.assertGreater(payload["operation_count"], 0)
        self.assertIn("full_sop", payload["unit_type_counts"])
        first_operation = payload["operations"][0]
        self.assertEqual(first_operation["operation"], "create_or_update_unit")
        self.assertTrue(first_operation["title"])
        self.assertIn(first_operation["review_status"], {"needs_review", "approved", "reviewed"})
        self.assertEqual(payload["source_coverage"]["source_element_count"], 3)

    def test_verification_report_blocks_conflicting_policy_units(self) -> None:
        chunks = [
            {
                "section": "full_sop",
                "heading": "Refund SOP",
                "content": "Refund policy.",
                "metadata": {
                    "unit_type": "full_sop",
                    "retrieval_scope": "document",
                    "review_status": "approved",
                    "extraction_status": "structured",
                    "source_refs": [{"source_type": "text", "line_start": 1}],
                },
            },
            {
                "section": "policy_rule",
                "heading": "Refund under 50k",
                "content": "Apply rounding.",
                "metadata": {
                    "unit_type": "policy_rule",
                    "service": "refund",
                    "case_type": "small_amount",
                    "rounding_applies": True,
                    "review_status": "approved",
                    "extraction_status": "structured",
                    "source_refs": [{"source_type": "text", "line_start": 2}],
                },
            },
            {
                "section": "policy_rule",
                "heading": "Refund under 50k exception",
                "content": "Do not apply rounding.",
                "metadata": {
                    "unit_type": "policy_rule",
                    "service": "refund",
                    "case_type": "small_amount",
                    "rounding_applies": False,
                    "review_status": "approved",
                    "extraction_status": "structured",
                    "source_refs": [{"source_type": "text", "line_start": 3}],
                },
            },
        ]

        report = verification_report_payload(chunks, "policy_rule")

        self.assertIn("conflicting_policy_units", report["hard_blockers"])
        self.assertEqual(report["reduce_reconcile"]["status"], "failed")
        self.assertEqual(report["reduce_reconcile"]["conflict_count"], 1)
        self.assertEqual(report["reduce_reconcile"]["conflicts"][0]["type"], "rounding_apply_conflict")

    def test_unsupported_semantic_unit_is_blocked(self) -> None:
        unit = SemanticUnit(
            unit_id="unit_missing",
            unit_type="generic_fact",
            fields={"content": "Unsupported claim"},
            source_element_ids=[],
            confidence=0.91,
        )

        result = validate_semantic_units_supported([unit])

        self.assertFalse(result.passed)
        self.assertIn("unit_missing:missing_source_elements", result.critical_warnings)

    def test_workflow_edge_requires_edge_level_evidence(self) -> None:
        graph = build_document_evidence_graph(
            filename="flow.pdf",
            content_type="application/pdf",
            data=b"%PDF",
            raw_text="",
            raw_context={
                "visual_layout": {
                    "pages": [
                        {
                            "page": 1,
                            "image_size": [1000, 1000],
                            "graph_candidate": {
                                "nodes": [
                                    {"id": "step_1", "text": "1 Start", "bbox": [1, 1, 10, 10]},
                                    {"id": "step_2", "text": "2 End", "bbox": [20, 20, 30, 30]},
                                ],
                                "edges": [
                                    {"from_node": "step_1", "to_node": "step_2", "condition": "next"},
                                ],
                            },
                        }
                    ]
                }
            },
            blocks=[],
            classification=None,
        )

        result = validate_workflow_graph(graph)

        self.assertFalse(result.passed)
        self.assertIn("workflow_edge_missing_source_ref", result.critical_warnings)

    def test_spreadsheet_chunks_cite_sheet_cell_locator(self) -> None:
        graph = build_document_evidence_graph(
            filename="rules.xlsx",
            content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
            data=b"fake",
            raw_text="# Rules\nA1: Refund policy",
            raw_context={"sheets": [("Rules", [(1, ["Policy", "Refund"]), (2, ["Action", "Review payment"])])]},
            blocks=[],
            classification=None,
        )
        chunks = legacy_chunks_from_evidence(graph, document_type="policy_table", source_type="spreadsheet")

        self.assertTrue(chunks)
        for chunk in chunks:
            refs = chunk.metadata.get("source_refs") or []
            self.assertTrue(any(ref.get("sheet") == "Rules" and ref.get("row_start") for ref in refs), chunk)
            self.assertTrue(chunk.metadata.get("source_element_ids"))
            self.assertTrue(chunk.metadata.get("evidence_hash"))

    def test_image_upload_creates_structured_source_element(self) -> None:
        graph = build_document_evidence_graph(
            filename="flow.png",
            content_type="image/png",
            data=b"\x89PNG\r\n",
            raw_text="Image asset uploaded: flow.png.",
            raw_context={},
            blocks=[],
            classification=None,
        )

        image_elements = [element for element in graph.source_elements if element.kind == "image_asset"]

        self.assertEqual(len(image_elements), 1)
        self.assertEqual(image_elements[0].element_id, "image_asset_1")
        self.assertEqual(image_elements[0].metadata["source_refs"][0]["source_type"], "image")
        self.assertEqual(image_elements[0].metadata["source_refs"][0]["mime_type"], "image/png")
        self.assertEqual(image_elements[0].provenance.extraction_method, "visual_asset")

    def test_image_upload_preserves_generated_caption_as_evidence(self) -> None:
        graph = build_document_evidence_graph(
            filename="flow.png",
            content_type="image/png",
            data=b"\x89PNG\r\n",
            raw_text="",
            raw_context={
                "image_captions": [
                    {
                        "text": "Flow chart showing refund escalation from CS to finance.",
                        "confidence": 0.82,
                        "model": "vision-caption-test",
                    }
                ]
            },
            blocks=[],
            classification=None,
        )

        caption_elements = [element for element in graph.source_elements if element.kind == "image_caption"]

        self.assertEqual(len(caption_elements), 1)
        self.assertEqual(caption_elements[0].text, "Flow chart showing refund escalation from CS to finance.")
        self.assertEqual(caption_elements[0].metadata["describes_element_id"], "image_asset_1")
        self.assertEqual(caption_elements[0].metadata["caption_status"], "generated")
        self.assertEqual(caption_elements[0].provenance.extraction_method, "visual_caption")
        self.assertTrue(
            any(relation.kind == "describes" and relation.source_id == "image_caption_1" and relation.target_id == "image_asset_1" for relation in graph.relations)
        )

    def test_extract_raw_evidence_adds_image_caption_context(self) -> None:
        from unittest.mock import patch

        with patch(
            "app.ingestion.describe_image_asset",
            return_value=(
                {"text": "Refund workflow screenshot with approval steps.", "confidence": 0.81, "model": "vision-test"},
                ["image_caption_generated"],
                "",
            ),
        ):
            raw_text, warnings, raw_context = extract_raw_evidence("flow.png", "image/png", b"\x89PNG\r\n")

        self.assertIn("Image asset uploaded", raw_text)
        self.assertIn("image_caption_generated", warnings)
        self.assertEqual(raw_context["image_captions"][0]["text"], "Refund workflow screenshot with approval steps.")
        self.assertEqual(raw_context["image_captions"][0]["model"], "vision-test")

    def test_embedded_images_create_caption_ready_source_elements(self) -> None:
        graph = build_document_evidence_graph(
            filename="policy.docx",
            content_type="application/vnd.openxmlformats-officedocument.wordprocessingml.document",
            data=b"docx",
            raw_text="Policy text",
            raw_context={
                "docx_blocks": [{"text": "Policy text", "paragraph_index": 0}],
                "embedded_images": [
                    {
                        "image_id": "rId5",
                        "page": 2,
                        "bbox": [10, 20, 300, 180],
                        "mime_type": "image/png",
                        "byte_size": 1234,
                        "caption": "Screenshot of the refund evidence form.",
                    }
                ],
            },
            blocks=[],
            classification=None,
        )

        image_elements = [element for element in graph.source_elements if element.kind == "image_asset"]
        caption_elements = [element for element in graph.source_elements if element.kind == "image_caption"]

        self.assertEqual(len(image_elements), 1)
        self.assertEqual(image_elements[0].element_id, "embedded_image_rid5")
        self.assertEqual(image_elements[0].page_number, 2)
        self.assertEqual(image_elements[0].bbox, [10.0, 20.0, 300.0, 180.0])
        self.assertTrue(image_elements[0].metadata["requires_caption"])
        self.assertEqual(image_elements[0].metadata["source_refs"][0]["source_type"], "embedded_image")
        self.assertEqual(caption_elements[0].text, "Screenshot of the refund evidence form.")
        self.assertEqual(caption_elements[0].metadata["describes_element_id"], "embedded_image_rid5")

    def test_docx_raw_evidence_extracts_embedded_images(self) -> None:
        import base64
        from io import BytesIO
        from tempfile import NamedTemporaryFile

        from docx import Document
        from app.text_processing import DOCX_CONTENT_TYPE

        tiny_png = base64.b64decode(
            "iVBORw0KGgoAAAANSUhEUgAAAAEAAAABCAQAAAC1HAwCAAAAC0lEQVR42mP8/x8AAwMCAO+/p9sAAAAASUVORK5CYII="
        )
        doc = Document()
        doc.add_paragraph("Policy text")
        with NamedTemporaryFile(suffix=".png") as image_file:
            image_file.write(tiny_png)
            image_file.flush()
            doc.add_picture(image_file.name)
            output = BytesIO()
            doc.save(output)

        _raw_text, _warnings, raw_context = extract_raw_evidence(
            "policy.docx",
            DOCX_CONTENT_TYPE,
            output.getvalue(),
        )

        self.assertEqual(len(raw_context["embedded_images"]), 1)
        self.assertEqual(raw_context["embedded_images"][0]["mime_type"], "image/png")
        self.assertGreater(raw_context["embedded_images"][0]["byte_size"], 0)
        self.assertIn("image", raw_context["embedded_images"][0]["image_id"])

    def test_docx_raw_evidence_captions_embedded_images_without_persisting_data_url(self) -> None:
        import base64
        from io import BytesIO
        from tempfile import NamedTemporaryFile
        from unittest.mock import patch

        from docx import Document
        from app.text_processing import DOCX_CONTENT_TYPE

        tiny_png = base64.b64decode(
            "iVBORw0KGgoAAAANSUhEUgAAAAEAAAABCAQAAAC1HAwCAAAAC0lEQVR42mP8/x8AAwMCAO+/p9sAAAAASUVORK5CYII="
        )
        doc = Document()
        doc.add_paragraph("Policy text")
        with NamedTemporaryFile(suffix=".png") as image_file:
            image_file.write(tiny_png)
            image_file.flush()
            doc.add_picture(image_file.name)
            output = BytesIO()
            doc.save(output)

        with patch(
            "app.ingestion.describe_image_asset",
            return_value=(
                {"text": "Embedded refund screenshot.", "confidence": 0.78, "model": "vision-test"},
                ["embedded_image_caption_generated"],
                "",
            ),
        ) as describe:
            _raw_text, warnings, raw_context = extract_raw_evidence(
                "policy.docx",
                DOCX_CONTENT_TYPE,
                output.getvalue(),
            )

        image = raw_context["embedded_images"][0]
        self.assertEqual(image["caption"], "Embedded refund screenshot.")
        self.assertEqual(image["caption_confidence"], 0.78)
        self.assertEqual(image["caption_model"], "vision-test")
        self.assertNotIn("image_data_url", image)
        self.assertIn("embedded_image_caption_generated", warnings)
        self.assertTrue(describe.call_args.kwargs["image_data_url"].startswith("data:image/png;base64,"))

    def test_markdown_evidence_preserves_heading_hierarchy_and_line_spans(self) -> None:
        raw_text = "# Policy\n\n## Refund\nCS checks payment.\n"
        blocks = parse_document_blocks("policy.md", "text/markdown", raw_text, {})
        graph = build_document_evidence_graph(
            filename="policy.md",
            content_type="text/markdown",
            data=raw_text.encode("utf-8"),
            raw_text=raw_text,
            raw_context={},
            blocks=blocks,
            classification=None,
        )

        markdown_elements = [element for element in graph.source_elements if element.metadata.get("ast_path")]
        self.assertTrue(markdown_elements)
        self.assertTrue(any(element.provenance.source_locator.get("line_start") == 3 for element in markdown_elements))
        self.assertTrue(any(element.metadata.get("section_path") == ["Policy", "Refund"] for element in markdown_elements))

    def test_chunk_support_validator_blocks_missing_evidence_hash(self) -> None:
        result = validate_chunks_supported(
            build_evidence_bound_chunks(
                graph=None,
                semantic_units=[
                    SemanticUnit(
                        unit_id="unit_1",
                        unit_type="generic_fact",
                        fields={"content": "No source"},
                        source_element_ids=[],
                    )
                ],
            )
        )

        self.assertFalse(result.passed)
        self.assertIn("unit_1:missing_source_elements", result.critical_warnings)


if __name__ == "__main__":
    unittest.main()
