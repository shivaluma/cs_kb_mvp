from __future__ import annotations

import unittest
from io import BytesIO

from docx import Document

from app import ingestion, openrouter
from app.openrouter import apply_reasoning_effort, normalize_source_evidence_formatter_payload, source_evidence_raw_fallback
from app.retrieval import to_result
from app.schemas import DisplayContext, DisplayHighlight, SourceAnchor
from app.text_processing import DOCX_CONTENT_TYPE, classify_document, extract_docx_structure


def test_mixed_docx_sop_creates_parent_table_rows_groups_and_titles() -> None:
    data = build_policy_docx_fixture()
    raw_text, blocks, tables = extract_docx_structure(data, filename="quy-dinh-phan-hoi.docx")
    classification = classify_document("quy-dinh-phan-hoi.docx", DOCX_CONTENT_TYPE, raw_text)

    chunks = ingestion.extract_mixed_docx_policy_chunks(
        "quy-dinh-phan-hoi.docx",
        DOCX_CONTENT_TYPE,
        raw_text,
        {"docx_blocks": blocks, "docx_tables": tables},
        classification,
    )
    normalized = ingestion.normalize_units(chunks)

    assert tables[0]["columns"] == ["Mail", "Macro", "Lưu ý"]

    full_sop = next(chunk for chunk in normalized if chunk.metadata["unit_type"] == "full_sop")
    assert full_sop.metadata["chunk_type"] == "full_sop"
    assert full_sop.metadata["retrieval_scope"] == "document"

    section_units = [chunk for chunk in normalized if chunk.metadata.get("chunk_type") == "parent_section"]
    assert section_units

    table_parent = next(chunk for chunk in normalized if chunk.metadata.get("chunk_type") == "parent_table")
    assert table_parent.metadata["unit_type"] == "macro_table"
    assert table_parent.metadata["table_id"] == "table_0"
    assert table_parent.metadata["parent_section_id"]

    row_chunks = [
        chunk
        for chunk in normalized
        if chunk.metadata.get("unit_type") == "macro_script"
        and chunk.metadata.get("source_ref_quality") == "table_row"
    ]
    assert len(row_chunks) >= 3
    open_row = next(chunk for chunk in row_chunks if "Open" in chunk.content or "Mở đầu" in chunk.content)
    assert open_row.metadata["parent_chunk_id"] == "macro_table_0"
    assert open_row.metadata["parent_section_id"]
    assert open_row.metadata["column_names"] == ["Mail", "Macro", "Lưu ý"]
    assert open_row.metadata["source_refs"][0]["table_index"] == 0
    assert open_row.metadata["source_refs"][0]["row_index"] == 1
    assert "Mail: Open / Mở đầu" in open_row.metadata["retrieval_text"]
    assert open_row.heading.startswith("EMAIL macro: Mail: Open / Mở đầu")

    handling_group = next(
        chunk
        for chunk in normalized
        if chunk.metadata.get("chunk_type") == "grouped_parent"
        and chunk.metadata.get("unit_type") == "handling_rule"
    )
    handling_children = [
        chunk
        for chunk in normalized
        if chunk.metadata.get("grouped_parent_unit_id") == handling_group.metadata["unit_id"]
    ]
    assert len(handling_children) == 2
    assert "chưa đề cập" in handling_group.content
    assert "đã đề cập" in handling_group.content

    wording_group = next(
        chunk
        for chunk in normalized
        if chunk.metadata.get("chunk_type") == "grouped_parent"
        and chunk.metadata.get("unit_type") == "wording_rule"
    )
    assert '"xin lỗi"' in wording_group.heading
    assert '"rất tiếc"' in wording_group.heading

    compliance_group = next(
        chunk
        for chunk in normalized
        if chunk.metadata.get("chunk_type") == "grouped_parent"
        and chunk.metadata.get("unit_type") == "compliance_rule"
    )
    compliance_children = [
        chunk
        for chunk in normalized
        if chunk.metadata.get("grouped_parent_unit_id") == compliance_group.metadata["unit_id"]
    ]
    assert len(compliance_children) >= 2
    assert any(child.metadata.get("risk_level") in {"high", "critical"} for child in compliance_children)

    titles = [chunk.heading for chunk in normalized if chunk.metadata.get("unit_type") != "full_sop"]
    assert len(titles) == len(set(titles))
    generic_titles = {"Quy định không tiết lộ", "Quy định wording xin lỗi", "Lưu ý vận hành"}
    assert not generic_titles.intersection(titles)
    assert any('Không nói "ĐÓNG HỖ TRỢ TẠI ĐÂY"' in title for title in titles)
    assert any('Khi nào dùng "xin lỗi"' in title for title in titles)
    assert any("quy trình xử lý nội bộ" in title.lower() for title in titles)
    assert any("[bộ phận chuyên môn]" in title for title in titles)


def test_docx_style_headings_override_numbered_heading_heuristic() -> None:
    doc = Document()
    doc.add_heading("Quy định phản hồi CS", level=1)
    doc.add_heading("1. Email", level=1)
    doc.add_paragraph("Mẫu câu email.")
    doc.add_heading("4. Khiếu nại đối tượng còn lại", level=1)
    doc.add_paragraph("Không cung cấp quy trình xử lý nội bộ.")
    buffer = BytesIO()
    doc.save(buffer)

    _raw_text, blocks, _tables = extract_docx_structure(buffer.getvalue(), filename="numbered-headings.docx")
    target = next(block for block in blocks if block.get("text") == "4. Khiếu nại đối tượng còn lại")

    assert target["section_path"] == ["4. Khiếu nại đối tượng còn lại"]


def test_formatter_rejects_array_summary_and_preserves_raw_fallback() -> None:
    parsed, warnings = normalize_source_evidence_formatter_payload(
        [{"title": "Email", "content": "Mail: Open / Mở đầu\nMacro: Xin chào anh/chị + Tên"}],
        document_title="Quy định phản hồi",
        raw_text="raw source",
    )

    assert parsed["title"] == "Quy định phản hồi"
    assert parsed["markdown"] == "raw source"
    assert parsed["sections"] == []
    assert parsed["coverage_report"]["normalization"] == "array_rejected_to_raw_text"
    assert "model_returned_array_rejected_to_raw_text" in warnings

    object_payload, object_warnings = normalize_source_evidence_formatter_payload(
        {"title": "Doc", "markdown": "Body", "sections": [], "warnings": [], "coverage_report": {}},
        document_title="Fallback",
        raw_text="raw source",
    )
    assert object_payload["markdown"] == "Body"
    assert object_warnings == []

    fallback = source_evidence_raw_fallback("source.docx", "raw source evidence", "invalid_json")
    assert fallback["markdown"] == "raw source evidence"
    assert "raw_source_evidence_preserved" in fallback["warnings"]


def test_formatter_extracts_div_wrapped_layout_sections_and_google_bboxes() -> None:
    parsed, warnings = normalize_source_evidence_formatter_payload(
        {
            "title": "Purchase order",
            "markdown": (
                '<div data-bbox="[100,50,180,950]" data-label="Title">Purchase order</div>\n'
                '<div data-bbox="[200,40,500,960]" data-label="Table">'
                '<table><tr><th rowspan="2">Item</th><th colspan="2">Qty</th></tr>'
                "<tr><th>Ordered</th><th>Backorder</th></tr><tr><td>030</td><td>4</td><td>1</td></tr></table>"
                "</div>"
            ),
            "sections": [],
            "warnings": [],
            "coverage_report": {},
        },
        document_title="Fallback",
        raw_text="Purchase order raw",
    )

    assert "source_evidence_sections_extracted_from_div_wrappers" in warnings
    assert parsed["format"] == "markdown_div_wrapped"
    assert parsed["bbox_order"] == "google_yxyx"
    assert parsed["sections"][0]["layout_label"] == "Title"
    assert parsed["sections"][0]["bbox"] == [50.0, 100.0, 950.0, 180.0]
    assert parsed["sections"][1]["layout_label"] == "Table"
    assert 'colspan="2"' in parsed["sections"][1]["markdown"]
    assert 'rowspan="2"' in parsed["sections"][1]["markdown"]


def test_source_evidence_formatter_sends_page_images_for_visual_pdfs(monkeypatch) -> None:
    captured: dict[str, object] = {}

    def fake_completion(payload: dict[str, object], _headers: dict[str, str]) -> str:
        captured["payload"] = payload
        return (
            '{"title":"Workflow","format":"markdown_div_wrapped","bbox_order":"google_yxyx",'
            '"markdown":"<div data-bbox=\\"[100,50,180,950]\\" data-label=\\"Text\\">1. KH/TX liên hệ Be qua Chat social</div>",'
            '"sections":[],"warnings":[],"coverage_report":{}}'
        )

    monkeypatch.setattr(openrouter, "enabled", lambda: True)
    monkeypatch.setattr(openrouter, "completion_content", fake_completion)
    output, warnings, error = openrouter.format_source_evidence_view(
        "workflow.pdf",
        "KH/TX liên h ệ Be qua Chat social",
        "workflow_diagram",
        "diagram_pdf",
        page_images=["data:image/jpeg;base64,abc"],
    )

    user_content = captured["payload"]["messages"][1]["content"]  # type: ignore[index]
    assert isinstance(user_content, list)
    assert user_content[0]["type"] == "text"
    assert user_content[1] == {"type": "image_url", "image_url": {"url": "data:image/jpeg;base64,abc"}}
    assert output["sections"][0]["bbox"] == [50.0, 100.0, 950.0, 180.0]
    assert "source_evidence_sections_extracted_from_div_wrappers" in warnings
    assert error == ""


def test_source_evidence_formatter_uses_parser_system_prompt(monkeypatch) -> None:
    captured: dict[str, object] = {}

    def fake_completion(payload: dict[str, object], _headers: dict[str, str]) -> str:
        captured["payload"] = payload
        return (
            '{"title":"Workflow","format":"markdown_div_wrapped","bbox_order":"google_yxyx",'
            '"markdown":"<div data-bbox=\\"[100,50,180,950]\\" data-label=\\"Text\\">1. KH/TX liên hệ Be</div>",'
            '"sections":[],"warnings":[],"coverage_report":{}}'
        )

    monkeypatch.setattr(openrouter, "enabled", lambda: True)
    monkeypatch.setattr(openrouter, "completion_content", fake_completion)
    openrouter.format_source_evidence_view(
        "workflow.pdf",
        "KH/TX liên h ệ Be",
        "workflow_diagram",
        "diagram_pdf",
        page_images=["data:image/jpeg;base64,abc"],
    )

    system_content = captured["payload"]["messages"][0]["content"]  # type: ignore[index]
    assert "You are a document parser" in system_content
    assert "bản nháp SOP" not in system_content


def test_reasoning_effort_is_added_to_openrouter_payloads() -> None:
    payload = {"model": "google/gemini-3.1-flash-lite-preview", "messages": []}

    apply_reasoning_effort(payload, "low")

    assert payload["reasoning"] == {"effort": "low"}


def test_retrieval_result_contract_contains_matched_parent_and_scroll_target() -> None:
    row = {
        "chunk_id": "chunk-row-1",
        "document_id": "doc-1",
        "version_id": "version-1",
        "title": "Quy định phản hồi",
        "source_filename": "source.docx",
        "version_number": 2,
        "chunk_index": 4,
        "section": "macro_script",
        "heading": "EMAIL macro: Mail: Open / Mở đầu",
        "content": "Mail: Open / Mở đầu. Macro: Xin chào anh/chị + Tên.",
        "score": 0.42,
        "lexical_score": 0.4,
        "vector_score": 0.2,
        "rank_source": ["lexical"],
        "metadata": {
            "unit_type": "macro_script",
            "chunk_type": "atomic_child",
            "parent_section_id": "section_email",
            "parent_chunk_id": "macro_table_0",
            "section_id": "section_email",
            "section_title": "Đối với Email",
            "section_path": ["Mẫu câu chào mở đầu", "Đối với Email"],
            "source_refs": [
                {
                    "source_type": "docx_table",
                    "source_file": "source.docx",
                    "table_index": 0,
                    "row_index": 1,
                    "column_names": ["Mail", "Macro", "Lưu ý"],
                }
            ],
        },
        "display_context": DisplayContext(
            display_unit_type="table_section",
            document_id="doc-1",
            document_title="Quy định phản hồi",
            section_id="section_email",
            section_title="Đối với Email",
            content="Mail: Open / Mở đầu. Macro: Xin chào anh/chị + Tên.",
            highlights=[
                DisplayHighlight(
                    chunk_id="chunk-row-1",
                    text="Mail: Open / Mở đầu. Macro: Xin chào anh/chị + Tên.",
                    match_strategy="table_row_anchor",
                    source_anchor=SourceAnchor(
                        sop_id="doc-1",
                        sop_version_id="version-1",
                        section_id="section_email",
                        block_id="table_0_row_1",
                        table_id="table_0",
                        row_index=1,
                    ),
                )
            ],
            source_anchor=SourceAnchor(
                sop_id="doc-1",
                sop_version_id="version-1",
                section_id="section_email",
                block_id="table_0_row_1",
                table_id="table_0",
                row_index=1,
            ),
        ),
    }

    result = to_result(row)

    assert result.matched_chunk.chunk_id == "chunk-row-1"
    assert result.matched_chunk.chunk_type == "atomic_child"
    assert result.parent.parent_section_id == "section_email"
    assert result.parent.parent_chunk_id == "macro_table_0"
    assert result.display.open_mode == "full_document"
    assert result.display.highlight_source_refs[0]["row_index"] == 1
    assert result.display.scroll_target.table_index == 0
    assert result.display.scroll_target.row_index == 1


def build_policy_docx_fixture() -> bytes:
    doc = Document()
    doc.add_heading("Quy định nội dung phản hồi tài xế, khách hàng", level=1)
    doc.add_heading("Mẫu câu chào mở đầu, chào kết và cách xưng hô", level=1)
    doc.add_heading("Đối với Email", level=2)
    table = doc.add_table(rows=1, cols=3)
    headers = table.rows[0].cells
    headers[0].text = "Mail"
    headers[1].text = "Macro"
    headers[2].text = "Lưu ý"
    for values in [
        ("Open / Mở đầu", "Xin chào anh/chị + Tên", "Đối tác bắt buộc đúng"),
        ("Body / Nội dung chính", "Anh/chị", ""),
        ("Close / Kết thúc", "Anh/chị + Tên", ""),
    ]:
        cells = table.add_row().cells
        for index, value in enumerate(values):
            cells[index].text = value
    doc.add_paragraph("Lưu ý: Đối với TX, CS phải kiểm tra họ tên và hình ảnh trên hệ thống trước khi phản hồi.")
    doc.add_heading("Đối với Call/Chat", level=2)
    doc.add_paragraph("Mở đầu: “Be xin chào, em là/em tên là xxx”.")
    doc.add_paragraph("Nếu KH/TX chưa đề cập vấn đề: CS có thể hỏi để xác định vấn đề cần hỗ trợ.")
    doc.add_paragraph("Nếu KH/TX đã đề cập vấn đề: CS không hỏi lại, chỉ xác nhận hoặc hỏi thêm nếu thông tin chưa rõ.")
    doc.add_heading('Quy tắc dùng từ "xin lỗi, rất tiếc"', level=1)
    doc.add_paragraph('Dùng "xin lỗi" khi vấn đề do be, TX hoặc nhà hàng gây ra.')
    doc.add_paragraph('Dùng "rất tiếc" cho các trường hợp còn lại.')
    doc.add_heading("Mẫu câu phản hồi đối với case trùng", level=1)
    doc.add_paragraph('Tuyệt đối không nói "ĐÓNG HỖ TRỢ TẠI ĐÂY".')
    doc.add_heading("KH/TX khiếu nại đối tượng còn lại", level=1)
    doc.add_paragraph("Không chủ động cung cấp ngưỡng chế tài hoặc lý do chế tài.")
    doc.add_paragraph("Không chủ động cung cấp quy trình xử lý nội bộ, SI call time, số lần gọi hoặc nội dung cuộc gọi.")
    doc.add_paragraph('Chỉ dùng cụm "[bộ phận chuyên môn]" trong case xem xét khóa tài khoản.')
    buffer = BytesIO()
    doc.save(buffer)
    return buffer.getvalue()


class SOPParentChildChunkingTest(unittest.TestCase):
    def test_mixed_docx_sop_creates_parent_table_rows_groups_and_titles(self) -> None:
        test_mixed_docx_sop_creates_parent_table_rows_groups_and_titles()

    def test_docx_style_headings_override_numbered_heading_heuristic(self) -> None:
        test_docx_style_headings_override_numbered_heading_heuristic()

    def test_formatter_rejects_array_summary_and_preserves_raw_fallback(self) -> None:
        test_formatter_rejects_array_summary_and_preserves_raw_fallback()

    def test_formatter_extracts_div_wrapped_layout_sections_and_google_bboxes(self) -> None:
        test_formatter_extracts_div_wrapped_layout_sections_and_google_bboxes()

    def test_reasoning_effort_is_added_to_openrouter_payloads(self) -> None:
        test_reasoning_effort_is_added_to_openrouter_payloads()

    def test_retrieval_result_contract_contains_matched_parent_and_scroll_target(self) -> None:
        test_retrieval_result_contract_contains_matched_parent_and_scroll_target()


if __name__ == "__main__":
    unittest.main()
