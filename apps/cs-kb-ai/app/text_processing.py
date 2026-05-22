from __future__ import annotations

import hashlib
import io
import base64
import re
import unicodedata
from dataclasses import dataclass, field
from typing import Any

from docx import Document as DocxDocument
from docx.oxml.table import CT_Tbl
from docx.oxml.text.paragraph import CT_P
from docx.table import Table as DocxTable
from docx.text.paragraph import Paragraph as DocxParagraph
from openpyxl import load_workbook
from pypdf import PdfReader

from app.config import settings


WORD_RE = re.compile(r"[\w]+", re.UNICODE)
HEADING_RE = re.compile(r"^\s*(#{1,6}\s+|[A-Z][A-Z0-9 _/-]{5,}:)\s*(.+?)\s*$")
BULLET_RE = re.compile(r"^\s*(?:[-*•‣▪]|\d+[.)]|[a-zA-Z][.)])\s+")
SENTENCE_END_RE = re.compile(r"(?<=[.!?。！？])\s+")
INCOMPLETE_CONNECTOR_RE = re.compile(
    r"(?:^|\s)(nếu|neu|thì|thi|đối với|doi voi|trường hợp|truong hop|bao gồm|bao gom|và|va|hoặc|hoac|or|and)\s*$",
    re.IGNORECASE,
)
CONDITION_RE = re.compile(r"(?:^|\s)(nếu|neu|trường hợp|truong hop|đối với|doi voi|khi|when|if)\b", re.IGNORECASE)
ACTION_RE = re.compile(
    r"(?:^|\s)(thì|thi|cần|can|phải|phai|xử lý|xu ly|chuyển|chuyen|kiểm tra|kiem tra|gửi|gui|tạo|tao|thực hiện|thuc hien|không được|khong duoc|được phép|duoc phep|must|should|do not)\b",
    re.IGNORECASE,
)
NOTE_RE = re.compile(r"^\s*(lưu ý|luu y|note|warning|cảnh báo|canh bao|script|sla|zt)\b", re.IGNORECASE)
URL_RE = re.compile(r"(?:https?://|www\.)[^\s)]+", re.IGNORECASE)
DOCX_CONTENT_TYPE = "application/vnd.openxmlformats-officedocument.wordprocessingml.document"
ROMAN_SECTION_RE = re.compile(r"^\s*([IVXLCDM]+)[.)]\s+(.+?)\s*$", re.IGNORECASE)
NUMBERED_SECTION_RE = re.compile(r"^\s*(\d{1,2})[.)]\s*(.+?)\s*$")
DOCX_TEXT_HEADING_RE = re.compile(r"^\s*(?:đối với|doi voi)\s+(.+?)\s*$", re.IGNORECASE)
DOCX_PROHIBITION_RE = re.compile(r"(tuyệt\s+đối\s+không|không\s+chủ\s+động\s+cung\s+cấp|quy\s+trình\s+xử\s+lý\s+nội\s+bộ|chế\s+tài|chấm\s+lỗi)", re.IGNORECASE)
INLINE_BULLET_MARKER_RE = re.compile(
    r"(^|\s)([-*•‣▪])\s+(?=(?:nếu|neu|kh|tx|cs|không|khong|chỉ|chi|trường hợp|truong hop)\b)",
    re.IGNORECASE,
)


@dataclass(frozen=True)
class Chunk:
    chunk_index: int
    section: str
    heading: str
    content: str
    token_count: int
    metadata: dict[str, Any] = field(default_factory=dict)


@dataclass(frozen=True)
class DocumentClassification:
    document_type: str
    source_type: str
    confidence: float
    requires_review: bool
    warnings: list[str] = field(default_factory=list)
    sub_type: str = ""
    structure_type: str = ""


def classify_document(filename: str, content_type: str, raw_text: str = "") -> DocumentClassification:
    lower_name = filename.lower()
    normalized = normalize_phrase(raw_text[:5000])

    if is_spreadsheet_file(lower_name, content_type):
        if looks_like_kb_index_workbook(raw_text):
            return DocumentClassification("kb_index_workbook", "excel_workbook", 0.94, True)
        return DocumentClassification("policy_table", "spreadsheet", 0.92, True)
    if is_image_file(lower_name, content_type):
        return DocumentClassification(
            "asset_sop",
            "image",
            0.72,
            True,
            ["image_ocr_or_vision_extraction_required"],
        )
    if lower_name.endswith(".pdf") and looks_like_workflow(normalized):
        return DocumentClassification("workflow_diagram", "diagram_pdf", 0.78, True)
    if lower_name.endswith(".docx") and looks_like_communication_guideline_docx(lower_name, normalized):
        return DocumentClassification(
            "policy_rule",
            "docx_policy_rule",
            0.88,
            True,
            ["effective_from_missing_needs_review", "mixed_docx_policy_detected"],
            sub_type="communication_guideline",
            structure_type="mixed_docx",
        )
    if looks_like_policy_rule(lower_name, normalized):
        return DocumentClassification(
            "policy_rule",
            "docx_policy_rule" if lower_name.endswith(".docx") else "text_policy_rule",
            0.84,
            True,
            ["effective_from_missing_needs_review"],
        )
    if any(term in normalized for term in ["macro", "script", "cau tra loi", "phan hoi mau"]):
        return DocumentClassification("macro_script", "text", 0.72, True)
    if lower_name.endswith((".pdf", ".docx", ".md", ".txt")):
        return DocumentClassification("text_sop", "text", 0.84, True)
    return DocumentClassification("unknown", "upload", 0.35, True, ["unknown_document_type"])


def extract_text(filename: str, content_type: str, data: bytes) -> tuple[str, list[str]]:
    warnings: list[str] = []
    lower_name = filename.lower()

    if len(data) > settings.max_upload_bytes:
        raise ValueError(f"file_too_large:{settings.max_upload_bytes}")

    if is_image_file(lower_name, content_type):
        text = (
            f"Image asset uploaded: {filename}. OCR/vision extraction is not configured in this local MVP. "
            "Keep this version in draft and curate the SOP manually before publishing."
        )
        warnings.append("image_ocr_or_vision_extraction_required")
    elif lower_name.endswith(".pdf") or content_type == "application/pdf":
        text = extract_pdf_text(data)
    elif is_spreadsheet_file(lower_name, content_type):
        text, spreadsheet_warnings, _ = extract_spreadsheet(lower_name, data)
        warnings.extend(spreadsheet_warnings)
    elif is_docx_file(lower_name, content_type):
        text = extract_docx_text(data)
    else:
        text = data.decode("utf-8", errors="ignore")

    text = normalize_whitespace(text)
    if not text:
        warnings.append("empty_text_after_extraction")
    return text, warnings


def is_spreadsheet_file(filename: str, content_type: str) -> bool:
    return filename.endswith((".xlsx", ".xlsm", ".xls")) or content_type in {
        "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        "application/vnd.ms-excel.sheet.macroEnabled.12",
        "application/vnd.ms-excel",
    }


def is_docx_file(filename: str, content_type: str) -> bool:
    return filename.endswith(".docx") or content_type == DOCX_CONTENT_TYPE


def is_image_file(filename: str, content_type: str) -> bool:
    return filename.endswith((".png", ".jpg", ".jpeg", ".webp", ".gif", ".tif", ".tiff")) or content_type.startswith("image/")


def looks_like_workflow(normalized_text: str) -> bool:
    workflow_terms = [
        "yes",
        "no",
        "co",
        "khong",
        "decision",
        "flow",
        "quy trinh",
        "xac minh",
        "neu",
        "truong hop",
    ]
    hits = sum(1 for term in workflow_terms if phrase_in_query(term, normalized_text))
    return hits >= 4


def looks_like_policy_rule(lower_name: str, normalized_text: str) -> bool:
    source = normalize_phrase(lower_name) + " " + normalized_text
    signals = [
        any(term in source for term in ["quy dinh", "policy", "rule", "nguyen tac", "dieu kien"]),
        any(term in source for term in ["neu", "thi", "truong hop", "if", "when", "condition"]),
        any(term in source for term in ["duoc phep", "khong duoc", "bat buoc", "can", "phai"]),
        any(term in source for term in ["xu ly", "chuyen", "kiem tra", "approve", "reject", "escalate"]),
        any(term in source for term in ["canh bao", "rui ro", "vi pham", "bao mat", "tuan thu", "compliance"]),
        any(term in source for term in ["=>", "->", "|"]),
    ]
    return sum(1 for hit in signals if hit) >= 3


def looks_like_communication_guideline_docx(lower_name: str, normalized_text: str) -> bool:
    source = normalize_phrase(lower_name) + " " + normalized_text
    has_policy_name = any(term in source for term in ["quy dinh", "noi dung phan hoi", "phan hoi tx kh", "mau cau", "giao tiep"])
    has_channel_mix = sum(1 for term in ["email", "call", "chat"] if term in source) >= 2
    has_audience = any(term in source for term in ["tx", "kh", "tai xe", "khach hang", "quy khach hang"])
    has_comm_content = any(term in source for term in ["macro", "script", "mau cau", "xin chao", "xin loi", "phan hoi"])
    has_operational_risk = bool(DOCX_PROHIBITION_RE.search(source)) or any(term in source for term in ["khong cung cap", "noi bo", "zt"])
    return has_policy_name and has_channel_mix and has_audience and (has_comm_content or has_operational_risk)


def looks_like_kb_index_workbook(raw_text: str) -> bool:
    sheet_names = {
        normalize_phrase(match.group(1))
        for match in re.finditer(r"^#\s+(.+?)\s*$", raw_text or "", flags=re.MULTILINE)
    }
    expected = {
        "overal",
        "quy dinh lam viec ccu pcu",
        "quy dinh chung",
        "driver rider",
        "driver cleaner",
        "rider",
        "cleaner",
        "mcu",
        "link lam viec",
        "vip",
        "tinh nang san pham moi",
    }
    hits = sum(1 for name in sheet_names if name in expected)
    has_router_terms = any(name in sheet_names for name in {"driver rider", "rider", "mcu", "cleaner"})
    has_tool_sheet = "link lam viec" in sheet_names
    has_sop_index = any(name in sheet_names for name in {"quy dinh chung", "quy dinh lam viec ccu pcu"})
    return hits >= 4 and (has_tool_sheet or has_router_terms) and has_sop_index


def extract_pdf_text(data: bytes) -> str:
    reader = PdfReader(io.BytesIO(data))
    pages = []
    for index, page in enumerate(reader.pages):
        page_text = page.extract_text() or ""
        if page_text.strip():
            pages.append(f"\n\n[page {index + 1}]\n{page_text}")
    return "\n".join(pages)


def render_pdf_pages_as_data_urls(data: bytes, max_pages: int = 3, scale: float = 1.6) -> tuple[list[str], list[str]]:
    try:
        import pypdfium2 as pdfium
    except Exception:
        return [], ["pdf_vision_render_unavailable:pypdfium2_missing"]

    try:
        pdf = pdfium.PdfDocument(data)
        images: list[str] = []
        for page_index in range(min(len(pdf), max_pages)):
            page = pdf[page_index]
            bitmap = page.render(scale=scale)
            image = bitmap.to_pil()
            buffer = io.BytesIO()
            image.save(buffer, format="JPEG", quality=82, optimize=True)
            encoded = base64.b64encode(buffer.getvalue()).decode("ascii")
            images.append(f"data:image/jpeg;base64,{encoded}")
        if not images:
            return [], ["pdf_vision_render_empty"]
        return images, [f"pdf_vision_pages_rendered:{len(images)}"]
    except Exception as exc:
        return [], [f"pdf_vision_render_failed:{exc.__class__.__name__}"]


def render_pdf_page_jpeg(data: bytes, page_number: int = 1, scale: float = 1.8) -> bytes:
    import pypdfium2 as pdfium

    pdf = pdfium.PdfDocument(data)
    if page_number < 1 or page_number > len(pdf):
        raise IndexError("page_out_of_range")
    page = pdf[page_number - 1]
    bitmap = page.render(scale=scale)
    image = bitmap.to_pil()
    buffer = io.BytesIO()
    image.save(buffer, format="JPEG", quality=86, optimize=True)
    return buffer.getvalue()


def extract_docx_structure(data: bytes, filename: str = "") -> tuple[str, list[dict[str, Any]], list[dict[str, Any]]]:
    if len(data) > settings.max_upload_bytes:
        raise ValueError(f"file_too_large:{settings.max_upload_bytes}")

    doc = DocxDocument(io.BytesIO(data))
    blocks: list[dict[str, Any]] = []
    raw_lines: list[str] = []
    tables: list[dict[str, Any]] = []
    section_stack: list[dict[str, Any]] = []
    paragraph_index = 0
    table_index = 0
    current_list_group = ""

    for block_kind, body_item in iter_docx_body_blocks(doc):
        if block_kind == "paragraph":
            paragraph = body_item
            assert isinstance(paragraph, DocxParagraph)
            block, next_section_stack, current_list_group = docx_paragraph_block(
                paragraph=paragraph,
                paragraph_index=paragraph_index,
                block_index=len(blocks),
                section_stack=section_stack,
                current_list_group=current_list_group,
                source_file=filename,
            )
            paragraph_index += 1
            if not block:
                continue
            section_stack = next_section_stack
            paragraph_blocks = expand_docx_inline_bullet_block(block, start_block_index=len(blocks))
            blocks.extend(paragraph_blocks)
            raw_lines.extend(str(item.get("text") or "") for item in paragraph_blocks if str(item.get("text") or "").strip())
            continue

        if block_kind != "table":
            continue
        table = body_item
        assert isinstance(table, DocxTable)
        table_blocks, table_payload = docx_table_blocks(
            table=table,
            table_index=table_index,
            start_block_index=len(blocks),
            section_path=section_path_from_stack(section_stack),
            source_file=filename,
        )
        table_index += 1
        if not table_blocks:
            continue
        current_list_group = ""
        blocks.extend(table_blocks)
        raw_lines.extend(str(block.get("text") or "") for block in table_blocks if str(block.get("text") or "").strip())
        tables.append(table_payload)
    return normalize_whitespace("\n".join(raw_lines)), blocks, tables


def iter_docx_body_blocks(doc: Any) -> list[tuple[str, DocxParagraph | DocxTable]]:
    output: list[tuple[str, DocxParagraph | DocxTable]] = []
    body = doc.element.body
    for child in body.iterchildren():
        if isinstance(child, CT_P):
            output.append(("paragraph", DocxParagraph(child, doc)))
        elif isinstance(child, CT_Tbl):
            output.append(("table", DocxTable(child, doc)))
    return output


def docx_paragraph_block(
    *,
    paragraph: DocxParagraph,
    paragraph_index: int,
    block_index: int,
    section_stack: list[dict[str, Any]],
    current_list_group: str,
    source_file: str = "",
) -> tuple[dict[str, Any], list[dict[str, Any]], str]:
    text = normalize_cell_text(paragraph.text)
    if not text:
        return {}, section_stack, ""

    style = str(getattr(getattr(paragraph, "style", None), "name", "") or "")
    numbering = docx_paragraph_numbering(paragraph)
    heading = docx_section_heading(text, style, numbering)
    next_section_stack = section_stack
    next_section_path = section_path_from_stack(section_stack)
    block_type = "paragraph"
    list_group = ""

    if heading:
        next_section_stack = next_docx_section_stack(section_stack, heading["text"], int(heading["level"]))
        next_section_path = section_path_from_stack(next_section_stack)
        block_type = "heading"
        current_list_group = ""
    elif is_docx_list_item(text, style, numbering):
        block_type = "list_item"
        list_group = current_list_group or f"list_{block_index}"
        current_list_group = list_group
    else:
        current_list_group = ""

    source_ref = {
        "source_type": "docx",
        **({"source_file": source_file} if source_file else {}),
        "paragraph_index": paragraph_index,
        "heading_path": next_section_path,
    }
    block = {
        "block_id": f"p{paragraph_index}",
        "block_type": block_type,
        "type": block_type,
        "index": block_index,
        "text": text,
        "section_path": next_section_path,
        "paragraph_index": paragraph_index,
        "table_index": None,
        "row_index": None,
        "cells": {},
        "style": style,
        "numbering": numbering,
        "list_group": list_group,
        "source_ref": source_ref,
        "source_refs": [source_ref],
    }
    return block, next_section_stack, current_list_group


def docx_table_blocks(
    *,
    table: DocxTable,
    table_index: int,
    start_block_index: int,
    section_path: list[str],
    source_file: str = "",
) -> tuple[list[dict[str, Any]], dict[str, Any]]:
    blocks: list[dict[str, Any]] = []
    rows: list[dict[str, Any]] = []
    columns: list[str] = []

    for table_row in table.rows:
        cell_values = [normalize_docx_cell_text(cell.text) for cell in table_row.cells]
        if not any(cell_values):
            continue
        if not columns:
            columns = unique_docx_headers(cell_values)
            header_text = " | ".join(columns)
            source_ref = docx_block_table_source_ref(table_index, 0, columns, section_path, source_file=source_file)
            blocks.append(
                {
                    "block_id": f"t{table_index}_h",
                    "block_type": "docx_table_header",
                    "type": "docx_table_header",
                    "index": start_block_index + len(blocks),
                    "text": header_text,
                    "section_path": section_path,
                    "paragraph_index": None,
                    "table_index": table_index,
                    "row_index": 0,
                    "headers": columns,
                    "columns": columns,
                    "cells": {column: column for column in columns},
                    "cell_values": cell_values,
                    "style": "table",
                    "numbering": {},
                    "source_ref": source_ref,
                    "source_refs": [source_ref],
                }
            )
            continue

        row_index = len(rows) + 1
        values = {
            columns[index]: cell_values[index]
            for index in range(min(len(columns), len(cell_values)))
            if columns[index]
        }
        row_text = " | ".join(cell_preview(cell) for cell in cell_values if cell)
        source_ref = docx_block_table_source_ref(table_index, row_index, columns, section_path, row_text, source_file=source_file)
        row_payload = {
            "row_index": row_index,
            "headers": columns,
            "columns": columns,
            "cells": values,
            "cell_values": cell_values,
            "cell_text": row_text,
            "values": values,
            "text": row_text,
            "section_path": section_path,
            "source_ref": source_ref,
            "source_refs": [source_ref],
        }
        rows.append(row_payload)
        blocks.append(
            {
                "block_id": f"t{table_index}_r{row_index}",
                "block_type": "docx_table_row",
                "type": "docx_table_row",
                "index": start_block_index + len(blocks),
                "text": row_text,
                "section_path": section_path,
                "paragraph_index": None,
                "table_index": table_index,
                "row_index": row_index,
                "headers": columns,
                "columns": columns,
                "cells": values,
                "cell_values": cell_values,
                "cell_text": row_text,
                "values": values,
                "style": "table",
                "numbering": {},
                "source_ref": source_ref,
                "source_refs": [source_ref],
            }
        )

    if not columns:
        return [], {"table_index": table_index, "headers": [], "columns": [], "rows": [], "section_path": section_path}
    return blocks, {"table_index": table_index, "headers": columns, "columns": columns, "rows": rows, "section_path": section_path}


def docx_block_table_source_ref(table_index: int, row_index: int, columns: list[str], section_path: list[str], cell_text: str = "", source_file: str = "") -> dict[str, Any]:
    return {
        "source_type": "docx_table",
        **({"source_file": source_file} if source_file else {}),
        "table_index": table_index,
        "row_index": row_index,
        "column_names": columns,
        "heading_path": section_path,
        **({"cell_text": cell_text[:1000]} if cell_text else {}),
    }


def docx_paragraph_numbering(paragraph: DocxParagraph) -> dict[str, Any]:
    p_pr = getattr(paragraph._p, "pPr", None)
    num_pr = getattr(p_pr, "numPr", None) if p_pr is not None else None
    if num_pr is None:
        return {}
    num_id = getattr(getattr(num_pr, "numId", None), "val", None)
    ilvl = getattr(getattr(num_pr, "ilvl", None), "val", None)
    return {
        **({"num_id": str(num_id)} if num_id is not None else {}),
        **({"level": int(ilvl)} if ilvl is not None else {}),
    }


def is_docx_list_item(text: str, style: str, numbering: dict[str, Any]) -> bool:
    style_key = style.lower()
    if any(signal in style_key for signal in ["list", "bullet", "number"]):
        return True
    if numbering:
        return True
    return bool(BULLET_RE.match(text))


def docx_section_heading(text: str, style: str, numbering: dict[str, Any]) -> dict[str, Any]:
    stripped = text.strip()
    roman = ROMAN_SECTION_RE.match(stripped)
    if roman and len(stripped) <= 180:
        return {"level": 1, "text": stripped, "kind": "roman_section"}

    style_key = style.lower()
    if style_key.startswith("heading") and len(stripped) <= 180:
        level_match = re.search(r"(\d+)", style_key)
        level = int(level_match.group(1)) if level_match else 1
        return {"level": min(max(level, 1), 4), "text": stripped, "kind": "style_heading"}

    numbered = NUMBERED_SECTION_RE.match(stripped)
    if numbered and len(stripped) <= 160:
        remainder = numbered.group(2).strip()
        if DOCX_TEXT_HEADING_RE.match(remainder) or len(remainder.split()) <= 10:
            return {"level": 2, "text": stripped, "kind": "numbered_subsection"}

    if DOCX_TEXT_HEADING_RE.match(stripped) and len(stripped) <= 120:
        return {"level": 2, "text": stripped, "kind": "text_heading"}

    if looks_like_docx_top_level_heading_text(stripped, style, numbering):
        return {"level": 1, "text": stripped, "kind": "inferred_top_level_heading"}

    return {}


def looks_like_docx_top_level_heading_text(text: str, style: str, numbering: dict[str, Any]) -> bool:
    if len(text) > 180 or BULLET_RE.match(text):
        return False
    normalized = normalize_phrase(text)
    if DOCX_TEXT_HEADING_RE.match(text):
        return False
    top_level_prefixes = (
        "mau cau",
        "quy tac",
        "quy dinh",
        "truong hop",
        "noi dung",
        "cach xung ho",
        "kiem soat",
    )
    if not normalized.startswith(top_level_prefixes):
        return False
    style_key = style.lower()
    has_list_or_numbering_context = bool(numbering) or any(signal in style_key for signal in ["list", "number", "heading"])
    return bool(has_list_or_numbering_context or len(text.split()) <= 14)


def section_path_from_stack(section_stack: list[dict[str, Any]]) -> list[str]:
    return [str(item.get("text") or "").strip() for item in sorted(section_stack, key=lambda item: int(item.get("level") or 0)) if str(item.get("text") or "").strip()]


def next_docx_section_stack(section_stack: list[dict[str, Any]], heading: str, level: int) -> list[dict[str, Any]]:
    next_stack = [
        item
        for item in section_stack
        if int(item.get("level") or 0) < level
    ]
    next_stack.append({"level": max(level, 1), "text": heading})
    return next_stack


def next_docx_section_path(current_path: list[str], heading: str, level: int) -> list[str]:
    if level <= 1:
        return [heading]
    output = current_path[: level - 1]
    while len(output) < level - 1:
        output.append("")
    output = [item for item in output if item]
    return [*output, heading]


def split_docx_inline_bullets(text: str) -> tuple[str, list[str]]:
    markers = list(INLINE_BULLET_MARKER_RE.finditer(text))
    if not markers:
        return text.strip(), []
    title = text[: markers[0].start(2)].strip(" :-\n\r\t")
    items: list[str] = []
    for index, marker in enumerate(markers):
        start = marker.end()
        end = markers[index + 1].start(2) if index + 1 < len(markers) else len(text)
        item = text[start:end].strip(" ;\n\r\t")
        if item:
            items.append(item)
    return title, items


def expand_docx_inline_bullet_block(block: dict[str, Any], *, start_block_index: int) -> list[dict[str, Any]]:
    block_type = str(block.get("block_type") or block.get("type") or "")
    if block_type == "heading" or block_type.startswith("docx_table"):
        return [block]
    title, items = split_docx_inline_bullets(str(block.get("text") or ""))
    if len(items) < 2:
        return [block]

    list_group = str(block.get("list_group") or f"list_{start_block_index}")
    output: list[dict[str, Any]] = []
    if title:
        parent = dict(block)
        parent_refs = [dict(ref, inline_group=True) for ref in block.get("source_refs", []) if isinstance(ref, dict)]
        parent["block_id"] = f"{block.get('block_id')}_group"
        parent["block_type"] = "list_group"
        parent["type"] = "list_group"
        parent["index"] = start_block_index
        parent["text"] = title
        parent["list_group"] = list_group
        parent["items"] = items
        if parent_refs:
            parent["source_refs"] = parent_refs
            parent["source_ref"] = parent_refs[0]
        output.append(parent)

    for item_index, item in enumerate(items, start=1):
        item_block = dict(block)
        item_refs = [
            dict(ref, inline_item_index=item_index)
            for ref in block.get("source_refs", [])
            if isinstance(ref, dict)
        ]
        item_block["block_id"] = f"{block.get('block_id')}_i{item_index}"
        item_block["block_type"] = "list_item"
        item_block["type"] = "list_item"
        item_block["index"] = start_block_index + len(output)
        item_block["text"] = item
        item_block["list_group"] = list_group
        item_block["numbering"] = {
            **(block.get("numbering") if isinstance(block.get("numbering"), dict) else {}),
            "inline_bullet": True,
            "inline_item_index": item_index,
        }
        if item_refs:
            item_block["source_refs"] = item_refs
            item_block["source_ref"] = item_refs[0]
        output.append(item_block)

    return output


def extract_docx_text(data: bytes) -> str:
    raw_text, _blocks, _tables = extract_docx_structure(data)
    return raw_text


def normalize_docx_cell_text(value: str) -> str:
    lines = [normalize_cell_text(line) for line in str(value or "").splitlines()]
    return "\n".join(line for line in lines if line)


def cell_preview(value: str) -> str:
    return re.sub(r"\s*\n\s*", " / ", value).strip()


def unique_docx_headers(values: list[str]) -> list[str]:
    output: list[str] = []
    seen: dict[str, int] = {}
    for index, value in enumerate(values, start=1):
        header = cell_preview(value) or f"Column {index}"
        count = seen.get(header, 0) + 1
        seen[header] = count
        output.append(header if count == 1 else f"{header} {count}")
    return output


def extract_spreadsheet(filename: str, data: bytes) -> tuple[str, list[str], list[Chunk]]:
    warnings: list[str] = []
    raw_blocks: list[str] = []
    chunks: list[Chunk] = []
    sheets = spreadsheet_rows(filename, data)

    for sheet_name, rows in sheets:
        section = slugify(sheet_name)[:80] or "sheet"
        sheet_chunks = spreadsheet_sheet_chunks(rows, section, sheet_name, len(chunks))
        if not sheet_chunks:
            warnings.append(f"sheet_empty:{sheet_name}")
            continue
        raw_blocks.append(f"# {sheet_name}")
        for chunk in sheet_chunks:
            source_refs = chunk.metadata.get("source_refs") if isinstance(chunk.metadata, dict) else []
            if source_refs and isinstance(source_refs[0], dict):
                source_refs[0]["source_file"] = filename
            chunk = Chunk(
                chunk_index=chunk.chunk_index,
                section=chunk.section,
                heading=chunk.heading,
                content=chunk.content,
                token_count=chunk.token_count,
                metadata={**chunk.metadata, "source_filename": filename, "source_refs": source_refs},
            )
            raw_blocks.append(f"## {chunk.heading}\n{chunk.content}")
            chunks.append(chunk)

    raw_text = "\n\n".join(raw_blocks)
    if not chunks:
        warnings.append("spreadsheet_no_indexable_rows")
    return raw_text, warnings, reindex_chunks(chunks)


def spreadsheet_rows(filename: str, data: bytes) -> list[tuple[str, list[tuple[int, list[str]]]]]:
    if filename.endswith(".xls") and not filename.endswith((".xlsx", ".xlsm")):
        import xlrd

        workbook = xlrd.open_workbook(file_contents=data)
        sheets = []
        for sheet in workbook.sheets():
            rows = []
            for row_index in range(sheet.nrows):
                values = [cell_to_text(sheet.cell_value(row_index, col_index)) for col_index in range(sheet.ncols)]
                if any(values):
                    rows.append((row_index + 1, trim_trailing_empty(values)))
            sheets.append((sheet.name.strip(), rows))
        return sheets

    workbook = load_workbook(io.BytesIO(data), read_only=False, data_only=True)
    sheets = []
    for worksheet in workbook.worksheets:
        rows = []
        for row_number, row in enumerate(worksheet.iter_rows(values_only=False), start=1):
            values = [cell_to_text_with_hyperlink(cell) for cell in row]
            if any(values):
                rows.append((row_number, trim_trailing_empty(values)))
        sheets.append((worksheet.title.strip(), rows))
    return sheets


def spreadsheet_sheet_chunks(non_empty_rows: list[tuple[int, list[str]]], section: str, sheet_name: str, start_index: int) -> list[Chunk]:
    chunks: list[Chunk] = []
    context: list[str] = []
    header_candidates: list[tuple[int, list[str]]] = []
    seen_data = False

    for row_number, values in non_empty_rows:
        filled_count = sum(1 for value in values if value)
        if filled_count == 0:
            continue

        if filled_count == 1:
            text = next(value for value in values if value)
            if not seen_data and len(text) <= 140:
                context.append(text)
                continue
            chunks.append(spreadsheet_chunk(start_index + len(chunks), section, sheet_name, row_number, context, "", text, values))
            seen_data = True
            continue

        if not seen_data and is_header_like(values):
            header_candidates.append((row_number, values))
            continue

        header_rows = [candidate_values for _, candidate_values in header_candidates]
        headers = flatten_headers(header_rows, len(values))
        chunks.extend(spreadsheet_row_chunks(start_index + len(chunks), section, sheet_name, row_number, context, values, headers))
        seen_data = True

    if not chunks and len(header_candidates) > 1:
        base_header = header_candidates[0][1]
        for row_number, values in header_candidates[1:]:
            headers = flatten_headers([base_header], len(values))
            chunks.extend(spreadsheet_row_chunks(start_index + len(chunks), section, sheet_name, row_number, context, values, headers))

    if not chunks and context:
        content = "\n".join(context)
        chunks.append(spreadsheet_chunk(start_index, section, sheet_name, 1, [], sheet_name, content, context))

    return chunks


def spreadsheet_row_chunks(
    index: int,
    section: str,
    sheet_name: str,
    row_number: int,
    context: list[str],
    values: list[str],
    headers: list[str],
) -> list[Chunk]:
    content = row_to_text(values, headers)
    heading = row_heading(values, headers, context)
    rule_id = f"{slugify(sheet_name) or 'sheet'}_row_{row_number}"
    max_atomic_tokens = max(settings.chunk_target_tokens * 4, 1000)
    if len(tokenize(content)) <= max_atomic_tokens:
        return [spreadsheet_chunk(index, section, sheet_name, row_number, context, heading, content, values, headers)]

    key_lines = []
    for column_index, value in enumerate(values):
        if value and len(value) <= 100:
            header = headers[column_index] if column_index < len(headers) else f"Column {column_index + 1}"
            key_lines.append(f"{header}: {value}")
        if len(key_lines) >= 2:
            break

    chunks = []
    for column_index, value in enumerate(values):
        if not value:
            continue
        header = headers[column_index] if column_index < len(headers) else f"Column {column_index + 1}"
        if len(value) <= 100 and key_lines:
            continue
        parts = split_text_by_tokens(value, settings.chunk_target_tokens)
        for part_index, part in enumerate(parts):
            suffix = f" part {part_index + 1}" if len(parts) > 1 else ""
            cell_content = "\n".join([*key_lines, f"{header}: {part}"]).strip()
            chunks.append(
                spreadsheet_chunk(
                    index + len(chunks),
                    section,
                    sheet_name,
                    row_number,
                    context,
                    f"{heading} / {header}{suffix}" if heading else f"{header}{suffix}",
                    cell_content,
                    values,
                    headers,
                    {"rule_id": rule_id, "parent_rule_id": rule_id, "split_from_row": True, "split_part": part_index + 1},
                )
            )
    return chunks or [spreadsheet_chunk(index, section, sheet_name, row_number, context, heading, content, values, headers)]


def split_text_by_tokens(value: str, target_tokens: int) -> list[str]:
    if len(tokenize(value)) <= target_tokens:
        return [value]
    sentences = split_sentences(value)
    parts: list[str] = []
    current: list[str] = []
    for sentence in sentences:
        candidate = " ".join([*current, sentence]).strip()
        if current and len(tokenize(candidate)) > target_tokens:
            parts.append(" ".join(current).strip())
            current = [sentence]
        else:
            current.append(sentence)
    if current:
        parts.append(" ".join(current).strip())
    return repair_incomplete_text_parts(parts)


def spreadsheet_chunk(
    index: int,
    section: str,
    sheet_name: str,
    row_number: int,
    context: list[str],
    heading: str,
    content: str,
    values: list[str] | None = None,
    headers: list[str] | None = None,
    extra_metadata: dict[str, Any] | None = None,
) -> Chunk:
    context_lines = [f"Ngữ cảnh: {item}" for item in context[-3:]]
    row_sentence = spreadsheet_row_sentence(sheet_name, row_number, values or [], headers or [])
    body = "\n".join([*context_lines, row_sentence, content]).strip()
    resolved_heading = heading or (context[-1] if context else f"{sheet_name} dòng {row_number}")
    rule_id = f"{slugify(sheet_name) or 'sheet'}_row_{row_number}"
    related_documents = spreadsheet_related_documents_from_row(values or [], headers or [], resolved_heading)
    hyperlinks = spreadsheet_hyperlinks_from_row(values or [], headers or [])
    related_metadata = {"related_documents": related_documents} if related_documents else {}
    return Chunk(
        chunk_index=index,
        section=section,
        heading=resolved_heading[:180],
        content=body,
        token_count=len(tokenize(body)),
        metadata={
            "sheet_name": sheet_name,
            "row_number": row_number,
            "row_index": row_number,
            "headers": headers or [],
            "row_values": values or [],
            "hyperlinks": hyperlinks,
            "source_type": "spreadsheet",
            "unit_type": "table_row",
            "rule_id": rule_id,
            "parent_unit_id": rule_id,
            "section_path": [sheet_name],
            "section_id": slugify(sheet_name) or "sheet",
            "section_title": sheet_name,
            "table_id": f"sheet_{slugify(sheet_name) or 'sheet'}",
            "block_id": f"sheet_{slugify(sheet_name) or 'sheet'}_row_{row_number}",
            "source_refs": [{"source_type": "excel", "source_file": "", "sheet": sheet_name, "row_start": row_number, "row_end": row_number, "column_names": headers or []}],
            **related_metadata,
            **(extra_metadata or {}),
        },
    )


def spreadsheet_row_sentence(sheet_name: str, row_number: int, values: list[str], headers: list[str]) -> str:
    pairs = []
    for index, value in enumerate(values):
        if not value:
            continue
        header = headers[index] if index < len(headers) and headers[index] else f"Column {index + 1}"
        pairs.append(f"{header}: {value}")
    if not pairs:
        return ""
    return f"Trong phần {sheet_name}, dòng {row_number} ghi " + "; ".join(pairs) + "."


def row_to_text(values: list[str], headers: list[str]) -> str:
    pairs = []
    for index, value in enumerate(values):
        if not value:
            continue
        header = headers[index] if index < len(headers) and headers[index] else f"Column {index + 1}"
        pairs.append(f"{header}: {value}")
    return "\n".join(pairs)


def row_heading(values: list[str], headers: list[str], context: list[str]) -> str:
    for index, value in enumerate(values):
        if value and len(value) <= 80:
            header = headers[index] if index < len(headers) else ""
            if header:
                return f"{header}: {value}"
            return value
    return context[-1] if context else ""


def flatten_headers(header_rows: list[list[str]], width: int) -> list[str]:
    if not header_rows:
        return [f"Column {index + 1}" for index in range(width)]

    filled_rows = [fill_horizontal(row, width) for row in header_rows]
    output = []
    for column in range(width):
        parts = []
        for row in filled_rows:
            if column < len(row) and row[column] and row[column] not in parts:
                parts.append(row[column])
        output.append(" / ".join(parts) if parts else f"Column {column + 1}")
    return output


def fill_horizontal(values: list[str], width: int) -> list[str]:
    output = values[:width] + [""] * max(0, width - len(values))
    last = ""
    for index, value in enumerate(output):
        if value:
            last = value
        elif last:
            output[index] = last
    return output


def is_header_like(values: list[str]) -> bool:
    non_empty = [value for value in values if value]
    if not non_empty:
        return False
    if any(URL_RE.search(value) for value in non_empty):
        return False
    if any(len(value) > 120 for value in non_empty):
        return False
    return True


def cell_to_text(value: Any) -> str:
    if value is None:
        return ""
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return normalize_cell_text(str(value))


def cell_to_text_with_hyperlink(cell: Any) -> str:
    text = cell_to_text(getattr(cell, "value", None))
    hyperlink = getattr(cell, "hyperlink", None)
    if not hyperlink:
        return text
    target = normalize_cell_text(str(getattr(hyperlink, "target", "") or getattr(hyperlink, "location", "") or ""))
    display = cell_to_text(getattr(hyperlink, "display", None))
    label = display or text or "Link"
    if target and target not in label and target not in text:
        return f"{label} ({target})"
    return label or target


def spreadsheet_related_documents_from_row(values: list[str], headers: list[str], fallback_title: str = "") -> list[dict[str, Any]]:
    context_title = first_non_link_value(values, headers) or fallback_title
    output: list[dict[str, Any]] = []
    seen: set[tuple[str, str]] = set()
    for index, value in enumerate(values):
        urls = URL_RE.findall(value or "")
        if not urls:
            continue
        header = headers[index] if index < len(headers) and headers[index] else f"Column {index + 1}"
        label = link_label_from_value(value)
        if is_generic_link_label(label):
            label = context_title or header
        for url in urls:
            key = (normalize_cell_text(label).lower(), url)
            if not label or key in seen:
                continue
            seen.add(key)
            output.append(
                {
                    "target_title": label[:300],
                    "relation_type": "references",
                    "source_url": url,
                    "source_header": header,
                    "source_value": value[:500],
                    "relation_source": "spreadsheet_hyperlink",
                }
            )
    return output


def spreadsheet_hyperlinks_from_row(values: list[str], headers: list[str]) -> list[dict[str, Any]]:
    output: list[dict[str, Any]] = []
    seen: set[str] = set()
    for index, value in enumerate(values):
        for url in URL_RE.findall(value or ""):
            if url in seen:
                continue
            seen.add(url)
            header = headers[index] if index < len(headers) and headers[index] else f"Column {index + 1}"
            label = link_label_from_value(value) or header
            output.append(
                {
                    "url": url,
                    "label": label[:300],
                    "column_name": header,
                    "cell_text": value[:1000],
                }
            )
    return output


def link_label_from_value(value: str) -> str:
    label = URL_RE.sub("", value or "")
    label = re.sub(r"\(\s*\)", "", label)
    return normalize_cell_text(label.strip(" -:()"))


def first_non_link_value(values: list[str], headers: list[str]) -> str:
    for index, value in enumerate(values):
        if not value or URL_RE.search(value):
            continue
        header = headers[index] if index < len(headers) else ""
        if is_generic_link_label(value) or is_generic_link_label(header):
            continue
        return value[:300]
    return ""


def is_generic_link_label(value: str) -> bool:
    normalized = normalize_phrase(value)
    return normalized in {
        "",
        "link",
        "url",
        "hyperlink",
        "xem tai day",
        "tai day",
        "link quy dinh",
        "duong dan",
        "link tai lieu",
        "tai lieu lien quan",
    }


def normalize_cell_text(value: str) -> str:
    return re.sub(r"\s+", " ", value.replace("\x00", " ")).strip()


def trim_trailing_empty(values: list[str]) -> list[str]:
    output = list(values)
    while output and not output[-1]:
        output.pop()
    return output


def reindex_chunks(chunks: list[Chunk]) -> list[Chunk]:
    return [
        Chunk(
            chunk_index=index,
            section=chunk.section,
            heading=chunk.heading,
            content=chunk.content,
            token_count=chunk.token_count,
            metadata=chunk.metadata,
        )
        for index, chunk in enumerate(chunks)
    ]


def ensure_full_sop_layer(
    chunks: list[Chunk],
    raw_text: str,
    filename: str,
    source_type: str,
    document_type: str,
) -> list[Chunk]:
    has_document_layer = any(
        chunk.section == "full_sop"
        or chunk.metadata.get("unit_type") == "full_sop"
        or chunk.metadata.get("retrieval_scope") == "document"
        for chunk in chunks
    )
    if has_document_layer:
        return reindex_chunks(chunks)

    content = normalize_cell_text(raw_text).strip()
    if not content:
        return reindex_chunks(chunks)
    full_content = content[:30000]
    if len(content) > len(full_content):
        full_content += "\n\n[Nội dung gốc dài hơn 30.000 ký tự và đã được rút gọn trong full SOP preview. Atomic units vẫn dùng các chunk bên dưới.]"

    document_chunk = Chunk(
        chunk_index=0,
        section="full_sop",
        heading=path_safe_title(filename),
        content=full_content,
        token_count=len(tokenize(full_content)),
        metadata={
            "unit_type": "full_sop",
            "retrieval_scope": "document",
            "source_type": source_type,
            "document_type": document_type,
            "source_filename": filename,
            "confidence": 0.72,
            "review_status": "needs_review",
            "source_refs": [default_source_ref(filename, source_type, full_content)],
            "source_ref_quality": default_source_ref_quality(filename),
            "production_ready_source_refs": default_source_ref_quality(filename) != "page_only",
            "source_ref_acknowledged": False if default_source_ref_quality(filename) == "page_only" else True,
        },
    )
    return reindex_chunks([document_chunk, *chunks])


def path_safe_title(filename: str) -> str:
    title = filename.rsplit("/", 1)[-1].rsplit("\\", 1)[-1]
    return title.rsplit(".", 1)[0].replace("_", " ").replace("-", " ").strip()[:180] or "Full SOP"


def default_source_ref(filename: str, source_type: str, content: str) -> dict[str, Any]:
    lower = filename.lower()
    if lower.endswith((".xlsx", ".xlsm", ".xls")):
        return {"source_type": "excel", "source_file": filename, "sheet": "unknown", "row_start": 1, "row_end": 1}
    if lower.endswith(".pdf"):
        return {"source_type": "pdf_diagram" if source_type == "diagram_pdf" else "pdf", "source_file": filename, "page": 1, "bbox": []}
    if lower.endswith(".docx"):
        return {"source_type": "docx", "source_file": filename, "paragraph_index": 0}
    line_count = max(1, content.count("\n") + 1)
    return {"source_type": source_type or "text", "source_file": filename, "line_start": 1, "line_end": line_count}


def default_source_ref_quality(filename: str) -> str:
    return "page_only" if filename.lower().endswith(".pdf") else "structured"


def workflow_units_to_chunks(units: list[dict[str, Any]], raw_text: str, filename: str) -> list[Chunk]:
    chunks: list[Chunk] = []
    effective_from = extract_effective_date(raw_text)
    workflow_id = workflow_identifier(filename, raw_text)
    for index, unit in enumerate(units):
        content = normalize_cell_text(str(unit.get("content") or ""))
        if not content:
            continue
        unit_type = slugify(str(unit.get("unit_type") or "workflow_step")) or "workflow_step"
        metadata = unit.get("metadata") if isinstance(unit.get("metadata"), dict) else {}
        chunks.append(
            Chunk(
                chunk_index=len(chunks),
                section=unit_type,
                heading=(str(unit.get("title") or vietnamese_unit_title(unit_type)).strip() or vietnamese_unit_title(unit_type))[:180],
                content=content,
                token_count=len(tokenize(content)),
                metadata={
                    "unit_type": unit_type,
                    "workflow_id": workflow_id,
                    "source_type": "diagram_pdf",
                    "source_page": metadata.get("source_page") or metadata.get("page_number") or 1,
                    "page_number": metadata.get("page_number") or metadata.get("source_page") or 1,
                    "source_ref_quality": metadata.get("source_ref_quality") or "page_only",
                    "production_ready_source_refs": metadata.get("production_ready_source_refs") if "production_ready_source_refs" in metadata else False,
                    "source_ref_acknowledged": metadata.get("source_ref_acknowledged") is True,
                    "effective_from": metadata.get("effective_from") or effective_from,
                    "confidence": unit.get("confidence", metadata.get("confidence", 0.74)),
                    **metadata,
                    "retrieval_scope": metadata.get("retrieval_scope") or ("document" if unit_type == "full_sop" else "unit"),
                },
            )
        )
    return chunks


def ai_units_to_chunks(units: list[dict[str, Any]], filename: str, source_type: str, document_type: str) -> list[Chunk]:
    chunks: list[Chunk] = []
    for unit in units:
        content = normalize_cell_text(str(unit.get("content") or ""))
        if not content:
            continue
        metadata = unit.get("metadata") if isinstance(unit.get("metadata"), dict) else {}
        unit_type = slugify(str(unit.get("unit_type") or metadata.get("unit_type") or "text_section")) or "text_section"
        confidence = unit.get("confidence", metadata.get("confidence", 0.72))
        try:
            confidence_value = float(confidence)
        except (TypeError, ValueError):
            confidence_value = 0.72
        chunks.append(
            Chunk(
                chunk_index=len(chunks),
                section=unit_type,
                heading=(str(unit.get("title") or vietnamese_unit_title(unit_type)).strip() or vietnamese_unit_title(unit_type))[:180],
                content=content,
                token_count=len(tokenize(content)),
                metadata={
                    "unit_type": unit_type,
                    "source_type": source_type,
                    "document_type": document_type,
                    "source_filename": filename,
                    "confidence": max(0.0, min(confidence_value, 1.0)),
                    "retrieval_scope": metadata.get("retrieval_scope") or ("document" if unit_type == "full_sop" else "unit"),
                    **metadata,
                },
            )
        )
    return chunks


def vietnamese_unit_title(unit_type: str) -> str:
    titles = {
        "workflow_overview": "Tổng quan quy trình",
        "workflow_graph": "Workflow graph",
        "candidate_action": "Candidate action",
        "candidate_decision": "Candidate decision",
        "candidate_annotation": "Candidate annotation",
        "candidate_sla": "Candidate SLA",
        "candidate_audit_rule": "Candidate audit rule",
        "candidate_queue_rule": "Candidate queue rule",
        "verification_dependency": "Tài liệu xác minh liên quan",
        "workflow_step": "Bước xử lý",
        "decision_point": "Điểm quyết định",
        "decision_rule": "Quy tắc quyết định",
        "sla_rule": "Quy định SLA",
        "escalation_rule": "Quy định chuyển xử lý",
        "case_creation_rule": "Quy định tạo case",
        "handoff_rule": "Quy định handoff",
        "operational_instruction": "Hướng dẫn thao tác",
        "routing_rule": "Quy định routing",
        "exception_rule": "Trường hợp ngoại lệ",
        "threshold_rule": "Quy định ngưỡng",
        "macro_table": "Bảng macro phản hồi",
        "wording_rule": "Quy định wording",
        "macro_script": "Script phản hồi",
        "operational_note": "Lưu ý vận hành",
        "security_note": "Lưu ý bảo mật",
        "compliance_rule": "Quy định tuân thủ",
        "example": "Ví dụ",
        "related_document": "Tài liệu liên quan",
        "issue_router_unit": "Dòng điều hướng vấn đề",
        "quick_action_rule": "Hành động nhanh",
        "sop_reference": "Tham chiếu SOP",
        "tool_link": "Công cụ làm việc",
        "vip_overlay_rule": "Quy định VIP",
        "product_update_note": "Ghi chú tính năng mới",
    }
    return titles.get(unit_type, "Đơn vị trích xuất cần review")


def extract_effective_date(text: str) -> str:
    match = re.search(r"(?:Áp dụng từ|Ap dung tu)\s*(\d{1,2})[/-](\d{1,2})[/-](\d{4})", text, re.IGNORECASE)
    if not match:
        return ""
    day, month, year = match.groups()
    return f"{year}-{int(month):02d}-{int(day):02d}"


def workflow_identifier(filename: str, raw_text: str) -> str:
    return slugify(filename)[:80] or "workflow"

def chunk_text(text: str, target_tokens: int | None = None, overlap_tokens: int | None = None) -> list[Chunk]:
    target = target_tokens or settings.chunk_target_tokens
    overlap = overlap_tokens or settings.chunk_overlap_tokens
    units = repair_logical_units(parse_logical_units(text))
    chunks: list[Chunk] = []
    current: list[dict[str, Any]] = []
    current_tokens = 0

    def flush() -> None:
        nonlocal current, current_tokens
        if not current:
            return
        chunks.append(chunk_from_units(len(chunks), current))
        if overlap > 0 and current:
            overlap_units: list[dict[str, Any]] = []
            overlap_count = 0
            for unit in reversed(current):
                unit_tokens = int(unit.get("token_count") or len(tokenize(str(unit.get("text") or ""))))
                if overlap_units and overlap_count + unit_tokens > overlap:
                    break
                overlap_units.insert(0, unit)
                overlap_count += unit_tokens
            current = overlap_units if overlap_units != current else []
            current_tokens = sum(int(unit.get("token_count") or 0) for unit in current)
        else:
            current = []
            current_tokens = 0

    for unit in units:
        unit_tokens = int(unit.get("token_count") or len(tokenize(str(unit.get("text") or ""))))
        if unit_tokens > target:
            if current:
                flush()
            for split_unit in split_large_logical_unit(unit, target):
                chunks.append(chunk_from_units(len(chunks), [split_unit]))
            current = []
            current_tokens = 0
            continue
        if current and current_tokens + unit_tokens > target:
            flush()
        current.append(unit)
        current_tokens += unit_tokens

    flush()
    return reindex_chunks(chunks)


def parse_logical_units(text: str) -> list[dict[str, Any]]:
    units: list[dict[str, Any]] = []
    section_path: list[str] = []
    pending_bullet: list[str] = []
    pending_bullet_start = 0

    def flush_bullet() -> None:
        nonlocal pending_bullet, pending_bullet_start
        if not pending_bullet:
            return
        content = "\n".join(pending_bullet).strip()
        units.append(logical_unit(content, "bullet", section_path, pending_bullet_start))
        pending_bullet = []
        pending_bullet_start = 0

    paragraphs = split_paragraphs(text)
    for index, paragraph in enumerate(paragraphs):
        heading = detect_heading(paragraph)
        if heading:
            flush_bullet()
            section_path = update_section_path(section_path, heading)
            continue
        if BULLET_RE.match(paragraph):
            if pending_bullet:
                flush_bullet()
            pending_bullet = [paragraph]
            pending_bullet_start = index
            continue
        if pending_bullet and looks_like_bullet_continuation(paragraph):
            pending_bullet.append(paragraph)
            continue
        flush_bullet()
        for sentence in split_sentences(paragraph):
            units.append(logical_unit(sentence, "sentence", section_path, index))
    flush_bullet()
    return units


def logical_unit(text: str, unit_type: str, section_path: list[str], source_index: int) -> dict[str, Any]:
    normalized = text.strip()
    return {
        "text": normalized,
        "logical_type": "note" if is_note_text(normalized) else unit_type,
        "section_path": list(section_path),
        "source_index": source_index,
        "token_count": len(tokenize(normalized)),
    }


def update_section_path(section_path: list[str], heading: str) -> list[str]:
    normalized = heading.strip()
    if not normalized:
        return section_path
    if re.match(r"^\d+\.\d+", normalized):
        return [*section_path[:1], normalized]
    if re.match(r"^\d+", normalized):
        return [normalized]
    return [normalized]


def split_sentences(text: str) -> list[str]:
    stripped = text.strip()
    if not stripped:
        return []
    if BULLET_RE.match(stripped):
        return [stripped]
    parts = [part.strip() for part in SENTENCE_END_RE.split(stripped) if part.strip()]
    if len(parts) <= 1:
        return [stripped]
    return parts


def looks_like_bullet_continuation(paragraph: str) -> bool:
    if BULLET_RE.match(paragraph) or detect_heading(paragraph):
        return False
    return True


def is_note_text(text: str) -> bool:
    return bool(NOTE_RE.search(text))


def repair_logical_units(units: list[dict[str, Any]]) -> list[dict[str, Any]]:
    repaired: list[dict[str, Any]] = []
    last_rule_or_step_index: int | None = None
    index = 0
    while index < len(units):
        unit = dict(units[index])
        text = str(unit.get("text") or "")
        while index + 1 < len(units) and should_merge_with_next(text):
            index += 1
            next_unit = units[index]
            text = "\n".join([text, str(next_unit.get("text") or "")]).strip()
            unit["text"] = text
            unit["token_count"] = len(tokenize(text))
        if is_note_text(text):
            if last_rule_or_step_index is not None:
                parent_id = repaired[last_rule_or_step_index].setdefault("unit_id", logical_unit_id(repaired[last_rule_or_step_index], last_rule_or_step_index))
                unit["attached_to"] = parent_id
            else:
                unit["review_status"] = "needs_review"
                unit["attachment_status"] = "needs_review_no_parent"
        elif CONDITION_RE.search(normalize_phrase(text)) or unit.get("logical_type") in {"sentence", "bullet"}:
            last_rule_or_step_index = len(repaired)
        repaired.append(unit)
        index += 1
    return repaired


def should_merge_with_next(text: str) -> bool:
    normalized = normalize_phrase(text)
    return bool(INCOMPLETE_CONNECTOR_RE.search(normalized)) or (has_condition(text) and not has_action(text))


def has_condition(text: str) -> bool:
    return bool(CONDITION_RE.search(normalize_phrase(text)))


def has_action(text: str) -> bool:
    return bool(ACTION_RE.search(normalize_phrase(text)))


def split_large_logical_unit(unit: dict[str, Any], target_tokens: int) -> list[dict[str, Any]]:
    text = str(unit.get("text") or "")
    if unit.get("logical_type") == "bullet":
        return [unit]
    sentences = split_sentences(text)
    if len(sentences) <= 1:
        return [unit]
    output: list[dict[str, Any]] = []
    current: list[str] = []
    for sentence in sentences:
        candidate = " ".join([*current, sentence]).strip()
        if current and len(tokenize(candidate)) > target_tokens:
            output.append({**unit, "text": " ".join(current).strip(), "token_count": len(tokenize(" ".join(current)))})
            current = [sentence]
        else:
            current.append(sentence)
    if current:
        output.append({**unit, "text": " ".join(current).strip(), "token_count": len(tokenize(" ".join(current)))})
    return repair_logical_units(output)


def repair_incomplete_text_parts(parts: list[str]) -> list[str]:
    repaired: list[str] = []
    index = 0
    while index < len(parts):
        text = parts[index]
        while index + 1 < len(parts) and should_merge_with_next(text):
            index += 1
            text = f"{text} {parts[index]}".strip()
        repaired.append(text)
        index += 1
    return repaired


def chunk_from_units(index: int, units: list[dict[str, Any]]) -> Chunk:
    first = units[0]
    section_path = first.get("section_path") or []
    section = slugify(" / ".join(section_path))[:80] or "body"
    heading = " / ".join(section_path[-2:])[:180] if section_path else ""
    content = "\n".join(str(unit.get("text") or "") for unit in units).strip()
    parent_unit_id = str(first.get("attached_to") or first.get("unit_id") or logical_unit_id(first, index))
    unit_type = infer_chunk_unit_type(units)
    return Chunk(
        chunk_index=index,
        section=section,
        heading=heading,
        content=content,
        token_count=len(tokenize(content)),
        metadata={
            "document_title": "",
            "section_path": section_path,
            "parent_unit_id": parent_unit_id,
            "unit_type": unit_type,
            "logical_unit_types": [unit.get("logical_type") for unit in units],
            "attached_to": first.get("attached_to", ""),
            "review_status": first.get("review_status", ""),
            "attachment_status": first.get("attachment_status", ""),
            "overlap_strategy": "logical_unit",
        },
    )


def logical_unit_id(unit: dict[str, Any], index: int) -> str:
    section = slugify(" ".join(unit.get("section_path") or []))[:40] or "body"
    text = slugify(str(unit.get("text") or ""))[:48] or f"unit_{index}"
    return f"{section}_{text}".strip("_")


def infer_chunk_unit_type(units: list[dict[str, Any]]) -> str:
    if any(unit.get("logical_type") == "note" for unit in units):
        return "operational_note"
    if any(unit.get("logical_type") == "bullet" for unit in units):
        return "checklist"
    text = "\n".join(str(unit.get("text") or "") for unit in units)
    if has_condition(text):
        return "candidate_rule"
    return "text_section"


def checksum(data: bytes) -> str:
    return hashlib.sha256(data).hexdigest()


def normalize_query(query: str, synonym_groups: list[dict] | None = None) -> str:
    normalized, _, _ = expand_query(query, synonym_groups)
    return normalized


def expand_query(query: str, synonym_groups: list[dict] | None = None) -> tuple[str, list[str], list[dict[str, Any]]]:
    normalized = normalize_phrase(query)
    expansions: list[str] = []
    matched_groups: list[dict[str, Any]] = []

    for group in synonym_groups or []:
        synonym_type = str(group.get("synonym_type") or "")
        if synonym_type == "placeholder":
            continue

        canonical = normalize_phrase(str(group.get("canonical_key") or ""))
        terms = unique_phrases([normalize_phrase(str(term.get("term") if isinstance(term, dict) else term)) for term in group.get("terms", [])])
        matched_terms = [term for term in terms if phrase_in_query(term, normalized)]
        matched_term = bool(matched_terms)
        matched_canonical = bool(canonical and phrase_in_query(canonical, normalized))

        if synonym_type == "regular" and (matched_term or matched_canonical):
            expansions.extend([canonical, *terms])
            matched_groups.append(explain_synonym_match(group, matched_terms, matched_canonical))
        elif synonym_type in {"one_way", "typo_correction"} and matched_term:
            expansions.append(canonical)
            matched_groups.append(explain_synonym_match(group, matched_terms, matched_canonical))

    normalized_query = join_unique([normalized, *expansions])
    return normalized_query, unique_phrases(expansions), matched_groups


def explain_synonym_match(group: dict, matched_terms: list[str], matched_canonical: bool) -> dict[str, Any]:
    return {
        "group_id": str(group.get("id") or ""),
        "canonical_key": str(group.get("canonical_key") or ""),
        "synonym_type": str(group.get("synonym_type") or ""),
        "domain": str(group.get("domain") or ""),
        "audience": str(group.get("audience") or ""),
        "matched_terms": matched_terms,
        "matched_canonical": matched_canonical,
    }


def normalize_phrase(text: str) -> str:
    return " ".join(tokenize(text.replace("_", " ")))


def phrase_in_query(phrase: str, query: str) -> bool:
    if not phrase or not query:
        return False
    return f" {phrase} " in f" {query} "


def join_unique(values: list[str]) -> str:
    seen = set()
    output = []
    for value in values:
        normalized = normalize_phrase(value)
        if not normalized or normalized in seen:
            continue
        seen.add(normalized)
        output.append(normalized)
    return " ".join(output).strip()


def unique_phrases(values: list[str]) -> list[str]:
    seen = set()
    output = []
    for value in values:
        normalized = normalize_phrase(value)
        if not normalized or normalized in seen:
            continue
        seen.add(normalized)
        output.append(normalized)
    return output


def tokenize(text: str) -> list[str]:
    text = text.replace("_", " ")
    return [normalize_token(match.group(0)) for match in WORD_RE.finditer(text) if normalize_token(match.group(0))]


def tokenize_raw(text: str) -> list[str]:
    return [match.group(0) for match in WORD_RE.finditer(text)]


def normalize_token(token: str) -> str:
    token = strip_accents(unicodedata.normalize("NFKC", token).lower())
    token = token.replace("_", " ")
    token = re.sub(r"[^\w]+", "", token, flags=re.UNICODE)
    return token


def strip_accents(value: str) -> str:
    decomposed = unicodedata.normalize("NFKD", value)
    stripped = "".join(char for char in decomposed if not unicodedata.combining(char))
    return stripped.replace("đ", "d").replace("Đ", "D")


def normalize_whitespace(text: str) -> str:
    lines = [re.sub(r"\s+", " ", line).strip() for line in text.replace("\x00", "").splitlines()]
    return "\n".join(line for line in lines if line)


def split_paragraphs(text: str) -> list[str]:
    parts = re.split(r"\n{1,}", text)
    return [part.strip() for part in parts if part.strip()]


def detect_heading(paragraph: str) -> str:
    if len(paragraph) > 120:
        return ""
    if re.match(r"^\s*[0-9]+[\.)]\s+", paragraph):
        return ""
    match = HEADING_RE.match(paragraph)
    if match:
        return match.group(2).strip()
    if paragraph.endswith(":") and 8 <= len(paragraph) <= 90:
        return paragraph[:-1].strip()
    return ""


def slugify(value: str) -> str:
    tokens = tokenize(value)
    return "_".join(tokens)
